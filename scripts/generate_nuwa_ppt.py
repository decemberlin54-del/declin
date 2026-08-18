#!/usr/bin/env python3
"""Generate 8-slide PPT for 《女娲造人》from HTML design."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/第23课女娲造人教学演示8页.pptx")
OUT_EN = Path("/workspace/nuwa-creation-myth-8slides.pptx")

SW, SH = Inches(13.333), Inches(7.5)
PL, PT = Inches(60 / 96), Inches(40 / 96)
PR, PB = Inches(60 / 96), Inches(40 / 96)
IW, IH = SW - PL - PR, SH - PT - PB

BG = RGBColor(0xF4, 0xF1, 0xEA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x2B, 0x2D, 0x42)
MUTED = RGBColor(0x6C, 0x75, 0x7D)
TEAL = RGBColor(0x6B, 0x90, 0x80)
OCHRE = RGBColor(0xE0, 0x9F, 0x68)
BORDER = RGBColor(0xE2, 0xE2, 0xD9)


def font(run, size=16, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    s.fill.solid()
    s.fill.fore_color.rgb = BG
    s.line.fill.background()
    sp = slide.shapes._spTree
    sp.remove(s._element)
    sp.insert(2, s._element)


def header(slide, title, num, total=8):
    tb = slide.shapes.add_textbox(PL, PT, IW * 0.75, Inches(0.45))
    r = tb.text_frame.paragraphs[0].add_run()
    font(r, 28, True, TEAL)
    r.text = title
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PL, PT + Inches(0.5), IW, Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()
    nb = slide.shapes.add_textbox(PL + IW * 0.82, PT + Inches(0.05), Inches(1.2), Inches(0.4))
    r2 = nb.text_frame.paragraphs[0].add_run()
    font(r2, 16, True, MUTED)
    r2.text = f"{num:02d} / {total:02d}"


def card(slide, left, top, width, height, title, lines, title_size=20):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = BORDER
    sh.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.15), top + Inches(0.28), Pt(4), Inches(0.35))
    bar.fill.solid()
    bar.fill.fore_color.rgb = OCHRE
    bar.line.fill.background()
    tf = slide.shapes.add_textbox(left + Inches(0.28), top + Inches(0.18), width - Inches(0.4), height - Inches(0.3)).text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    font(r0, title_size, True, TEAL)
    r0.text = title
    p0.space_after = Pt(10)
    for line in lines:
        p = tf.add_paragraph()
        r = p.add_run()
        font(r, 16 if len(line) < 80 else 14)
        r.text = line
        p.space_after = Pt(6)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


def slide1_cover(prs):
    s = new_slide(prs)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, SW / 2 - Inches(2.5), PT + Inches(1.5), Inches(5), Inches(0.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = TEAL
    sh.line.fill.background()
    tb = s.shapes.add_textbox(SW / 2 - Inches(2.5), PT + Inches(1.55), Inches(5), Inches(0.4))
    r = tb.text_frame.paragraphs[0].add_run()
    font(r, 14, False, RGBColor(0xFF, 0xFF, 0xFF))
    r.text = "七年级语文上册第六单元 · 第23课"
    tb.text_frame.paragraphs[0].alignment = 1

    for y, txt, sz, c, bd in [
        (2.3, "《女娲造人》", 44, TEXT, True),
        (3.2, "远古的设想与生命的赞歌", 22, TEAL, False),
        (4.0, "改写：袁珂", 18, MUTED, False),
    ]:
        t = s.shapes.add_textbox(PL, Inches(y), IW, Inches(0.7))
        t.text_frame.paragraphs[0].alignment = 1
        r = t.text_frame.paragraphs[0].add_run()
        font(r, sz, bd, c)
        r.text = txt


def slide2_compare(prs):
    s = new_slide(prs)
    header(s, "古今文本对比：从梗概到神话", 2)
    cw = (IW - Inches(0.2)) / 2
    y = PT + Inches(0.75)
    h = Inches(4.8)
    card(s, PL, y, cw, h, "古籍《风俗通》原典", [
        "“俗说天地开辟，未有人民，女娲抟黄土作人。剧务，力不暇供，乃引绳于泥中，举以为人。”",
        "",
        "特点：语言极简，仅记录基本骨架，无情绪与细节描述。",
    ])
    card(s, PL + cw + Inches(0.2), y, cw, h, "袁珂改编的现代神话", [
        "• 增加了女娲池边孤独的心理描绘。",
        "• 细化了捏泥人、赋予生命的过程。",
        "• 融入了人类开口叫“妈妈”的温馨场景。",
        "",
        "本质：用合理的想象填补历史的空白。",
    ])


def slide3_dimensions(prs):
    s = new_slide(prs)
    header(s, "想象的 4 个扩写维度", 3)
    items = [
        ("1. 动机赋予", "因荒凉与孤独而萌生创造同伴的渴望。"),
        ("2. 过程细节", "揉团黄泥、池水映照、挥洒泥浆等动作描写。"),
        ("3. 人性注入", "展现女娲如母亲般慈爱、欣喜与疲惫的心情。"),
        ("4. 繁衍逻辑", "区分男女，设立婚姻，解决人类世代绵延问题。"),
    ]
    cw = (IW - Inches(0.45)) / 4
    y = PT + Inches(0.8)
    for i, (t, d) in enumerate(items):
        card(s, PL + i * (cw + Inches(0.15)), y, cw, Inches(4.5), t, [d], 16)


def slide4_reality(prs):
    s = new_slide(prs)
    header(s, "远古先民的现实依托", 4)
    items = [
        ("现象：人从何而来？", "现实依托：泥土是孕育万物的基础，且具有极高的可塑性。因此想象出“黄泥造人”。"),
        ("现象：为什么人有贫富？", "现实依托：手揉泥人精致，藤条洒泥粗糙。以此解释现实中的阶级差异。"),
        ("现象：人类如何繁衍？", "现实依托：观察到男女交配与婚姻制度，推演为女娲安排男女婚配。"),
    ]
    cw = (IW - Inches(0.3)) / 3
    y = PT + Inches(0.8)
    for i, (t, d) in enumerate(items):
        card(s, PL + i * (cw + Inches(0.15)), y, cw, Inches(4.8), t, [d])


def slide5_character(prs):
    s = new_slide(prs)
    header(s, "形象剖析：神性与人性的交融", 5)
    cw = (IW - Inches(0.2)) / 2
    y = PT + Inches(0.75)
    card(s, PL, y, cw, Inches(4.8), "神性维度（造物主）", [
        "• 神通广大、变幻莫测。",
        "• 具有创造生命的至高权力。",
        "• 挥绳洒泥即可化为活生生的人。",
    ])
    card(s, PL + cw + Inches(0.2), y, cw, Inches(4.8), "人性维度（伟大的母亲）", [
        "• 会感到寂寞与孤独。",
        "• 看到作品时充满成功与满足感。",
        "• 工作累了会疲倦。",
        "• 神性为骨，人性为血肉。",
    ])


def slide6_essence(prs):
    s = new_slide(prs)
    header(s, "神话思维的本质", 6)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PL, PT + Inches(0.85), IW, Inches(4.6))
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = BORDER
    t1 = s.shapes.add_textbox(PL + Inches(0.5), PT + Inches(1.5), IW - Inches(1), Inches(0.8))
    t1.text_frame.paragraphs[0].alignment = 1
    r = t1.text_frame.paragraphs[0].add_run()
    font(r, 24, True, TEAL)
    r.text = "神话 = 现实困惑 + 现实对应物 + 浪漫推演"
    t2 = s.shapes.add_textbox(PL + Inches(0.8), PT + Inches(2.5), IW - Inches(1.6), Inches(2.5))
    t2.text_frame.word_wrap = True
    r2 = t2.text_frame.paragraphs[0].add_run()
    font(r2, 18)
    r2.text = (
        "神话绝非胡言乱语，而是远古先民在科技极度不发达的时代，"
        "基于对自然与社会的客观观察，运用强大想象力对世界做出的浪漫解释。"
    )


def slide7_criteria(prs):
    s = new_slide(prs)
    header(s, "“好想象”的评分三要素", 7)
    items = [
        ("1. 有依托（40%）", "必须能在生活中找到影子，不能凭空捏造。例如：泥土的可塑性。"),
        ("2. 有逻辑（30%）", "因果链条严密。例如：因为孤独所以造人，因为太累所以改用藤条。"),
        ("3. 有新奇（30%）", "提供意料之外却情理之中的视角。例如：泥人落地即能喊“妈妈”。"),
    ]
    cw = (IW - Inches(0.3)) / 3
    y = PT + Inches(0.8)
    for i, (t, d) in enumerate(items):
        card(s, PL + i * (cw + Inches(0.15)), y, cw, Inches(4.8), t, [d])


def slide8_task(prs):
    s = new_slide(prs)
    header(s, "拓展与创意表达", 8)
    card(s, PL, PT + Inches(0.75), IW, Inches(4.8), "微型神话创作", [
        "选答题目：",
        "• 请选择一个自然现象（例如：为什么会有四季、太阳为什么东升西落、雷电是如何产生的）。",
        "",
        "任务要求：",
        "1. 假想你自己是一名远古先民。",
        "2. 找到一个现实依托，展开合理想象。",
        "3. 写一段 150 字左右的远古神话片段。",
    ])


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    slide1_cover(prs)
    slide2_compare(prs)
    slide3_dimensions(prs)
    slide4_reality(prs)
    slide5_character(prs)
    slide6_essence(prs)
    slide7_criteria(prs)
    slide8_task(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    prs.save(str(OUT_EN))
    print(f"✓ {OUT} ({len(prs.slides)} slides)")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    build()
