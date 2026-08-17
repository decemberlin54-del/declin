#!/usr/bin/env python3
"""Generate lesson 8 design: 综合研读·想象说明卡."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path("/workspace/第六单元课时设计/第8课时把想象变成作品综合研读想象说明卡.docx")
BLANK = " " * 60


def set_font(run, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(12)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def para(doc, text="", bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run(text)
    set_font(r, bold=bold)


def blanks(doc, n=2):
    for _ in range(n):
        para(doc, BLANK)


def table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for ci, h in enumerate(headers):
        tbl.rows[0].cells[ci].text = h
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            tbl.rows[ri].cells[ci].text = val
    return tbl


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)

    para(doc, "第八课时   把想象变成作品   综合研读：为文本制作想象说明卡", bold=True)
    para(doc, "【课时目标】")
    para(
        doc,
        "1.通过比较梳理四篇课文，归纳神话、童话、寓言、神魔小说中想象的依据、特点与表达效果，"
        "完成“想象比较表”，发展审美鉴赏与逻辑思维能力。",
    )
    para(
        doc,
        "2.运用“文本依据—想象依据—表现方式—现实意味”的方法，为所选课文精彩片段制作想象说明卡，"
        "做到有文本依据、表达清楚，发展语言建构与运用能力。",
    )
    para(
        doc,
        "3.联系“想象与真实”母题及四条策展主题（自由与规则、真实与虚假、创造与生命、智慧与局限），"
        "能用规范语言阐释课文想象对当今生活的启示，为“想象力博物馆”策展做好准备，"
        "对接课标“联想和想象”要求及中考情境化、主题探究考查方向。",
    )
    para(doc, "评价任务：")
    para(
        doc,
        "1.完成表格一《四体想象比较》，能从想象依据、表现方式、表达效果等角度归纳不同文体的想象规律。（检测目标1）",
    )
    para(
        doc,
        "2.独立完成一张“想象说明卡”，写清文本片段、想象依据、表现方式与现实意味，语言准确、条理清楚。（检测目标2）",
    )
    para(
        doc,
        "3.结合策展主题，用30—60秒语言阐释所选课文的想象如何照见“想象与真实”，并能联系生活谈启示。（检测目标3）",
    )
    para(doc, "学习过程：")
    para(doc, "课前热身：")
    para(
        doc,
        "回顾第2—7课时的学习：你在《小圣施威降大圣》里追踪过“七十二变对照表”，"
        "在《皇帝的新装》里分析过“看不见的布”，在《女娲造人》里感受过“妈妈”一声呼唤的温度，"
        "在《寓言四则》里读懂过“白送”“穿井得一人”的讽刺与警醒。"
        "今天，你要从“读者”变成“策展人”——为“想象力博物馆”制作第一张正式的“想象说明卡”。"
        "请先完成热身填空：（指向目标1）",
    )
    para(doc, "四篇课文对应的二级母题：")
    para(doc, "《小圣施威降大圣》—________    《皇帝的新装》—________")
    para(doc, "《女娲造人》—________          《寓言四则》—________")
    blanks(doc, 1)

    para(doc, "学习任务一：综合比较，归纳“想象”规律")
    para(
        doc,
        "1.自主回顾四篇课文，完成表格一《四体想象比较》。"
        "提示：可联系第2课时“变化链”、第3—4课时“夸张反讽”、"
        "第5课时“古籍改写”、第6—7课时“情节与寓意”等学习成果。（指向目标1）",
    )
    para(doc, "表格一：四体想象比较")
    table(
        doc,
        ["篇目", "文体", "想象依据（从哪来）", "表现方式", "表达效果/二级母题"],
        [
            ["《小圣施威降大圣》", "神魔小说", "", "变化斗法、相克逻辑", "自由与规则"],
            ["《皇帝的新装》", "童话", "", "夸张、反讽", "真实与虚假"],
            ["《女娲造人》", "神话", "", "改写、添情", "创造与生命"],
            ["《寓言四则》", "寓言", "", "拟人、夸张、反转", "智慧与局限"],
        ],
    )
    para(doc, "")
    para(doc, "2.根据表格信息，概括不同文体想象的共同点与不同点。（检测目标1）")
    para(doc, "我的发现：四篇课文的想象都源于________，但表现方式不同：神魔小说重________，童话重________，神话重________，寓言重________。")
    blanks(doc, 2)
    para(
        doc,
        "3.链接课标与中考：课标要求“发挥联想和想象，感受文学的奇思妙想”；"
        "中考常考“概括内容—分析手法—探究主题—联系现实”。"
        "请用一句话说明：本单元为什么要同时学习四种想象类文体？（检测目标1）",
    )
    blanks(doc, 2)

    para(doc, "学习任务二：研读范例，学会制作“想象说明卡”")
    para(
        doc,
        "1.阅读“范例说明卡”，把握制作要求。"
        "想象说明卡须回答四个问题：选了哪段文字？想象从哪来？怎样表现？照见了什么真实？（指向目标2）",
    )
    para(doc, "【范例说明卡】《小圣施威降大圣》")
    para(doc, "文本片段：孙悟空变作土地庙，“只有尾巴不好收拾，竖在后面，变做一根旗竿。”")
    para(doc, "想象依据：生活中庙宇有旗竿，但旗竿不会竖在庙后——源于生活观察，又故意留下破绽。")
    para(doc, "表现方式：连续变化、细节描写、相克逻辑（一方变，一方识破）。")
    para(doc, "现实意味：自由可贵，但变化须受环境、形态与对手判断的制约——对应策展主题“自由与规则”。")
    para(doc, "")
    para(
        doc,
        "2.从四篇课文中任选一篇（建议选你第2—7课时笔记最丰富的一篇），"
        "完成学习任务单二“我的想象说明卡”。（检测目标2）",
    )
    para(doc, "学习任务单二：我的想象说明卡")
    table(
        doc,
        ["项目", "内容（请结合具体课文语句填写）"],
        [
            ["所选课文", ""],
            ["文本片段（可摘原文）", ""],
            ["想象依据（它从什么现实经验/材料出发？）", ""],
            ["表现方式（变化/夸张/改写/拟人等）", ""],
            ["现实意味（照见怎样的生活道理？）", ""],
            ["对应策展主题", "□自由与规则  □真实与虚假  □创造与生命  □智慧与局限"],
        ],
    )
    para(doc, "")
    para(
        doc,
        "3.对照“中考主题探究”答题思路，检查你的说明卡是否做到："
        "有文本依据、有分析过程、有主题提升。（指向目标2）",
    )
    para(doc, "自检清单：□引用了课文原句  □说清了“想象从哪来”  □点明了“照见什么真实”  □联系了策展主题")
    blanks(doc, 1)

    para(doc, "学习任务三：小组互评，完善说明卡，准备策展")
    para(doc, "1.小组内交换说明卡，按评价标准互评，提出一条修改建议。（指向目标3）")
    para(doc, "“想象说明卡”评价标准")
    table(
        doc,
        ["评价要点", "具体内容", "自评", "互评"],
        [
            ["文本依据", "准确引用或概括课文片段，不脱离原文", "☆☆☆☆☆", "☆☆☆☆☆"],
            ["想象分析", "说清想象依据与表现方式，分析合理", "☆☆☆☆☆", "☆☆☆☆☆"],
            ["主题提升", "能联系“想象与真实”及策展主题谈启示", "☆☆☆☆☆", "☆☆☆☆☆"],
            ["语言表达", "条理清楚，用语准确，无歧义", "☆☆☆☆☆", "☆☆☆☆☆"],
        ],
    )
    para(doc, "")
    para(doc, "同学修改建议：")
    blanks(doc, 2)
    para(
        doc,
        "2.模拟“想象力博物馆”30秒讲解：面向同学介绍你的展品，"
        "先说文本片段，再说想象亮点，最后说现实启示。（检测目标3）",
    )
    para(doc, "讲解提纲：")
    para(doc, "①展品名称：________    ②想象亮点：________    ③现实启示：________")
    blanks(doc, 2)
    para(
        doc,
        "3.创新任务：若把四张说明卡按“自由→真实→创造→智慧”的顺序布展，"
        "你会怎样设计参观路线？说说这样安排的理由。（检测目标3）",
    )
    blanks(doc, 2)

    para(doc, "四、课后小结")
    para(doc, "1.你能说出四篇课文想象的共同点与不同点吗？")
    blanks(doc, 1)
    para(doc, "2.制作想象说明卡，对你理解“想象与真实”有什么帮助？")
    blanks(doc, 1)
    para(doc, "3.下一课时将进行联想与想象写作/课本剧创编，你的说明卡可以为创作提供什么准备？")
    blanks(doc, 1)

    para(doc, "五、作业与检测：")
    para(doc, "1.必做：根据互评建议，完善“想象说明卡”，誊写工整，准备下节课布展使用。")
    para(doc, "2.必做：从另外三篇课文中，各用一句话写“如果我也为它做一张说明卡，我会选哪个片段”。")
    blanks(doc, 2)
    para(
        doc,
        "3.选做：查阅一处中考真题（如2025年云南中考“综合性学习”“阅读探究”类试题），"
        "说说本课“想象说明卡”的学习方法，与中考“结合材料谈启示”的答题有何相似之处。",
    )
    blanks(doc, 2)

    doc.save(str(OUT))
    # English alias for easy download
    alias = Path("/workspace/unit6-lesson8-imagination-card-design.docx")
    import shutil
    shutil.copy2(OUT, alias)
    print(f"✓ {OUT}")
    print(f"✓ {alias}")


if __name__ == "__main__":
    build()
