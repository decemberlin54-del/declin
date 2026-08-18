#!/usr/bin/env python3
"""Generate 七年级上册语文期末考试双向细目表 (portrait, 8 columns)."""

from pathlib import Path

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("/workspace/七年级语文期末考试双向细目表.docx")
OUT_EN = Path("/workspace/grade7-final-exam-spec-table.docx")


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def style_run(run, name="宋体", size=9, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def create_specification_document(filename=OUT):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("七年级上册语文期末考试双向细目表")
    style_run(run_title, "黑体", 18, True, RGBColor(31, 78, 121))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("基于课标要求与教育测量学规范命题设计")
    style_run(run_sub, "楷体", 12, False, RGBColor(89, 89, 89), italic=True)

    doc.add_paragraph()

    headers = [
        "题号",
        "对应板块",
        "考查知识点 / 核心素养项",
        "题型",
        "分值",
        "能力层级",
        "预估难度",
        "对应课标与教材要求",
    ]

    data = [
        ["1", "语言积累与运用", "汉字字音辨析（含多音字、易错字）", "选择题", "2", "识记", "0.85", "随文识字，准确掌握常见汉字读音。"],
        ["2", "语言积累与运用", "汉字字形辨析（同音字、形近字）", "选择题", "2", "识记", "0.85", "正确书写常用规范汉字。"],
        ["3", "语言积累与运用", "词语（成语）语义与语境搭配", "选择题", "2", "表达应用", "0.70", "理解词语在具体语境中的含义与用法。"],
        ["4", "语言积累与运用", "病句辨析与修改（介词滥用、动宾搭配）", "选择题", "2", "表达应用", "0.65", "辨析并修改常见语法与逻辑语病。"],
        ["5", "语言积累与运用", "句意衔接与逻辑排序", "选择题", "2", "表达应用", "0.65", "掌握段落内部语句的连贯性与逻辑结构。"],
        ["6", "语言积累与运用", "古诗文名篇名句默写", "填空题", "6", "识记", "0.80", "积累优秀诗文，准确默写名句名篇。"],
        ["7", "综合性学习", "情境化表达与阅读观点思辨", "简答题", "2", "分析综合", "0.70", "基于文本材料进行辩证思考与情境化表达。"],
        ["8", "综合性学习", "名著主题迁移与阅读推荐", "简答题", "3", "分析综合", "0.65", "提炼名著核心特质，结合现实意义规范表达。"],
        ["9", "名著阅读", "字词释义与书名内涵/作者情感映射", "简答题", "2", "理解/分析综合", "0.60", "结合具体篇目理解名著书名内涵与情感。"],
        ["10", "名著阅读", "人物形象、具体情节与抽象主题综合分析", "简答题", "3", "分析综合", "0.60", "理解名著人物形象，将情节上升为抽象精神。"],
        ["11", "古诗文阅读", "诗歌内容理解与艺术手法赏析判断", "选择题", "2", "分析综合", "0.70", "理解诗歌基本意象、情感及表达技巧。"],
        ["12", "古诗文阅读", "诗歌超越时空想象手法的比较分析", "简答题", "2", "分析综合", "0.55", "比较分析不同诗歌在抒情机制上的构思特征。"],
        ["13", "古诗文阅读", "文言实词/虚词含义解释", "填空题", "4", "理解", "0.65", "掌握常见文言实词在特定语境中的含义。"],
        ["14", "古诗文阅读", "文言文重点语句翻译", "简答题", "4", "理解", "0.60", "准确翻译文言文关键句（落实关键词语）。"],
        ["15", "古诗文阅读", "跨文本（甲乙文）内容比较与成因概括", "简答题", "2", "分析综合", "0.60", "筛选归纳文言文本关键信息并进行对比。"],
        ["16", "古诗文阅读", "角色迁移与主旨开导性表达", "简答题", "3", "表达应用/鉴赏评价", "0.55", "深入理解文章主旨，完成角色语意化劝告。"],
        ["17", "现代文阅读", "说明文信息提取、概括与说明方法作用", "填空/简答", "8", "理解/分析综合", "0.65", "准确提取文本关键信息，分析说明文语言特征。"],
        ["18", "现代文阅读", "说明文段落结构功能与主旨拓展作用分析", "简答题", "3", "分析综合", "0.55", "剖析补充段落在文章结构与深化主题上的作用。"],
        ["19", "现代文阅读", "跨文体（说明文与生态随笔）写作目的异同对比", "简答题", "3", "鉴赏评价", "0.50", "比较不同文体在表达目的与抒情视角上的差异。"],
        ["20", "现代文阅读", "散文中人物心理与情感变化的概括", "简答题", "2", "分析综合", "0.65", "捕捉文本关键情节，准确概括人物心理变化。"],
        ["21", "现代文阅读", '核心象征意象（"伞"）的表层与深层含义解析', "简答题", "3", "理解/鉴赏评价", "0.60", "解析散文中意象的深层内涵与哲理情感。"],
        ["22", "现代文阅读", "跨文本意象修辞与表达效果对比分析", "简答题", "3", "鉴赏评价", "0.55", "对比不同文章在修辞手法的运用与表达效果上的异同。"],
        ["23", "现代文阅读", "多文本主题整合与个人认知感悟表达", "简答题", "4", "鉴赏评价", "0.50", "结合多篇阅读材料，综合归纳主题并阐述个人感悟。"],
        ["24", "写作", "命题/半命题作文（细节特写、心理转折、主题升华）", "写作题", "40", "表达应用", "0.70", "写作记叙文或议论文，做到内容充实、结构清晰。"],
    ]

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [
        Inches(0.5),
        Inches(1.1),
        Inches(1.8),
        Inches(0.7),
        Inches(0.5),
        Inches(1.0),
        Inches(0.6),
        Inches(1.8),
    ]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=120, bottom=120)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            style_run(run, "黑体", 9.5, True, RGBColor(255, 255, 255))

    center_cols = {0, 3, 4, 5, 6}
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F2F4F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if c_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in p.runs:
                style_run(run, "宋体", 9)

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    doc.save(str(filename))
    print(f"✓ {filename}")


def main():
    create_specification_document(OUT)
    import shutil

    shutil.copy2(OUT, OUT_EN)
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    main()
