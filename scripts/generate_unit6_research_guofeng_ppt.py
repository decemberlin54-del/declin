#!/usr/bin/env python3
"""教研汇报 PPT：新中式极简国风 · 5 slides."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/单元学历案教研汇报国风.pptx")
OUT_EN = Path("/workspace/unit6-lesson-plan-research-guofeng.pptx")

SW = Inches(13.333)
SH = Inches(7.5)
PL, PT = Inches(0.72), Inches(0.42)
IW = SW - Inches(1.44)

# 新中式配色：宣纸白 · 黛蓝 · 朱砂 · 赭石 · 竹青
BG = RGBColor(0xF5, 0xF0, 0xE6)          # 宣纸白
DAILAN = RGBColor(0x3D, 0x5A, 0x73)      # 黛蓝
DAILAN_WASH = RGBColor(0xE4, 0xEA, 0xEF)
RED = RGBColor(0xC8, 0x3C, 0x23)         # 朱砂红
OCHRE = RGBColor(0x9B, 0x6B, 0x4A)       # 赭石
BAMBOO = RGBColor(0x4A, 0x7C, 0x6F)      # 竹青
BAMBOO_WASH = RGBColor(0xE6, 0xED, 0xE9)
INK = RGBColor(0x1E, 0x1E, 0x1E)
MUTED = RGBColor(0x5A, 0x6A, 0x72)

FONT_TITLE = "KaiTi"
FONT_BODY = "Microsoft YaHei"
TOTAL = 5


def font(run, size=16, bold=False, color=INK, title=False):
    run.font.name = FONT_TITLE if title else FONT_BODY
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    sh.fill.solid()
    sh.fill.fore_color.rgb = BG
    sh.line.fill.background()
    tree = slide.shapes._spTree
    tree.remove(sh._element)
    tree.insert(2, sh._element)


def wash(slide, l, t, w, h, color=DAILAN_WASH):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def line_box(slide, l, t, w, h, fill=BG, border=DAILAN, lw=Pt(1.2)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = lw
    return sh


def cloud_div(slide, y, w=None):
    w = w or IW
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PL, y, w, Inches(0.01))
    ln.fill.solid()
    ln.fill.fore_color.rgb = MUTED
    ln.line.fill.background()
    tb(slide, PL, y - Inches(0.07), w, Inches(0.22), [("～ ～ ～ ～ ～ ～ ～ ～ ～ ～ ～ ～", 9, False, MUTED)], PP_ALIGN.CENTER)


def seal(slide, l, t, size, text, fs=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, size, size)
    sh.fill.background()
    sh.line.color.rgb = RED
    sh.line.width = Pt(2)
    tb(slide, l, t, size, size, [(text, fs, True, RED)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def bar(slide, l, t, h, color=DAILAN, w=Inches(0.05)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, title=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(item, tuple):
            txt, sz, b, c = item
            r = p.add_run()
            font(r, sz, b, c, title=title or (b and i == 0))
            r.text = txt
        else:
            r = p.add_run()
            font(r, title=title)
            r.text = item
        p.space_after = Pt(4)
    return box


def hdr(slide, main, sub, n):
    seal(slide, PL, PT, Inches(0.42), f"{n:02d}", 10)
    tb(slide, PL + Inches(0.55), PT, IW * 0.72, Inches(0.42), [(main, 22, True, INK)], title=True)
    tb(slide, PL + Inches(0.55), PT + Inches(0.42), IW * 0.72, Inches(0.32), [(sub, 12, False, OCHRE)])
    tb(slide, PL + IW * 0.78, PT, IW * 0.22, Inches(0.35), [(f"卷 {n}/{TOTAL}", 11, False, MUTED)], PP_ALIGN.RIGHT)
    cloud_div(slide, PT + Inches(0.82))


def verse_card(slide, l, t, w, h, lines, accent=BAMBOO, fill=BAMBOO_WASH):
    wash(slide, l, t, w, h, fill)
    line_box(slide, l, t, w, h, fill, accent, Pt(1))
    bar(slide, l, t, h, accent)
    tb(slide, l + Inches(0.18), t + Inches(0.14), w - Inches(0.28), h - Inches(0.22), lines)


# ── Page 1: 封面 ──
def s01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    # 水墨云气与书页隐现
    wash(s, Inches(8.2), Inches(0.35), Inches(4.6), Inches(2.8), DAILAN_WASH)
    wash(s, Inches(9.5), Inches(2.2), Inches(3.4), Inches(2.2), BAMBOO_WASH)
    wash(s, Inches(0.4), Inches(5.2), Inches(3.8), Inches(1.85), DAILAN_WASH)
    # 书页意象
    line_box(s, Inches(9.8), Inches(1.0), Inches(2.6), Inches(3.4), BG, OCHRE, Pt(0.8))
    wash(s, Inches(10.0), Inches(1.15), Inches(2.2), Inches(3.0), BG)
    tb(
        s,
        Inches(10.15),
        Inches(1.5),
        Inches(1.9),
        Inches(2.2),
        [
            ("书", 28, True, MUTED),
            ("页", 28, True, MUTED),
            ("隐", 28, True, MUTED),
            ("现", 28, True, MUTED),
        ],
        PP_ALIGN.CENTER,
    )
    seal(s, PL, PT, Inches(0.48), "六", 11)
    tb(
        s,
        PL,
        Inches(1.55),
        Inches(8.2),
        Inches(1.6),
        [("在想象的光影里，", 38, True, INK)],
        title=True,
    )
    tb(
        s,
        PL,
        Inches(2.55),
        Inches(8.2),
        Inches(1.0),
        [("照见真实的自己", 38, True, DAILAN)],
        title=True,
    )
    bar(s, PL, Inches(3.75), Inches(0.65), RED, Inches(3.2))
    tb(
        s,
        PL + Inches(0.12),
        Inches(3.82),
        Inches(7.5),
        Inches(0.55),
        [("基于核心素养的七年级上册第六单元学历案设计与实践", 15, False, MUTED)],
    )
    cloud_div(s, Inches(5.0), Inches(8.5))
    tb(s, PL, Inches(5.25), Inches(5.0), Inches(0.4), [("汇报人 · 朱琳", 15, False, OCHRE)])
    tb(
        s,
        PL,
        Inches(6.55),
        IW,
        Inches(0.35),
        [("一切想象都有来处，奇幻故事照见生活本色", 13, False, BAMBOO)],
        PP_ALIGN.CENTER,
    )


# ── Page 2: 立意·课程重构 ──
def s02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    hdr(s, "立意 · 课程重构", "幻境寻真：在奇思妙想中认出生活本色", 2)
    line_box(s, PL, Inches(1.0), IW, Inches(1.05), DAILAN_WASH, DAILAN, Pt(1.5))
    tb(
        s,
        PL + Inches(0.25),
        Inches(1.12),
        IW - Inches(0.5),
        Inches(0.85),
        [
            ("核心母题", 12, True, RED),
            ("一切想象都有来处，奇幻故事是现实生活的延伸与隐喻", 16, True, INK),
            ("大任务驱动：「想象力博物馆」策展人（9课时递进游廊）", 13, False, OCHRE),
        ],
    )
    texts = [
        ("《小圣施威降大圣》", "变化之局", "辨析自由与规则"),
        ("《皇帝的新装》", "虚妄之镜", "洞察真实与虚假"),
        ("《女娲造人》", "生灵之爱", "感悟创造与生命"),
        ("《寓言四则》", "人世之鉴", "审视智慧与局限"),
    ]
    cw = (IW - Inches(0.18)) / 2
    ch = Inches(1.55)
    for i, (work, mood, theme) in enumerate(texts):
        col, row = i % 2, i // 2
        x = PL + col * (cw + Inches(0.18))
        y = Inches(2.35) + row * (ch + Inches(0.15))
        verse_card(
            s,
            x,
            y,
            cw,
            ch,
            [
                (work, 14, True, DAILAN),
                (mood, 13, True, RED),
                (theme, 13, False, INK),
            ],
            BAMBOO if i % 2 else DAILAN,
            BAMBOO_WASH if i % 2 else DAILAN_WASH,
        )
    line_box(s, PL, Inches(5.55), IW, Inches(0.95), BG, OCHRE, Pt(1))
    tb(
        s,
        PL + Inches(0.22),
        Inches(5.68),
        IW - Inches(0.44),
        Inches(0.72),
        [
            ("理论支撑", 12, True, RED),
            ("UbD 逆向设计（以终为始）  ·  情境学习理论（真实角色驱动）", 14, False, INK),
        ],
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )


# ── Page 3: 实践·课堂探索 ──
def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    hdr(s, "实践 · 课堂探索", "架桥通冥：以理性的支架托举飞扬的思维", 3)
    wash(s, PL, Inches(1.0), IW, Inches(0.72), BAMBOO_WASH)
    line_box(s, PL, Inches(1.0), IW, Inches(0.72), BAMBOO_WASH, BAMBOO, Pt(1))
    tb(
        s,
        PL + Inches(0.22),
        Inches(1.12),
        IW - Inches(0.44),
        Inches(0.52),
        [("教学痛点：避开「看热闹」的浮光掠影，走向「悟门道」的理性思维", 14, True, INK)],
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )
    scaffolds = [
        ("风云互变 · 推理链", "探寻《降大圣》中因果相克逻辑，悟「自由须受因果限制」"),
        ("荒诞隐喻 · 心理图", "剖析《新装》中的集体说谎，探「真话困境」的现实土壤"),
        ("温情注脚 · 情感谱", "对比《女娲造人》古籍原典，体会声声「妈妈」的生命温度"),
        ("意象迁移 · 联想桥", "从「伞」的四维联想（形/用/情/境），助推合情合理之想象"),
    ]
    y = Inches(2.0)
    for i, (title, desc) in enumerate(scaffolds):
        accent = DAILAN if i % 2 == 0 else BAMBOO
        fill = DAILAN_WASH if i % 2 == 0 else BAMBOO_WASH
        seal(s, PL, y + Inches(0.08), Inches(0.38), f"{i + 1}", 10)
        verse_card(
            s,
            PL + Inches(0.52),
            y,
            IW - Inches(0.52),
            Inches(0.98),
            [(title, 14, True, accent), (desc, 12, False, INK)],
            accent,
            fill,
        )
        y += Inches(1.12)
    tb(
        s,
        PL,
        Inches(6.55),
        IW,
        Inches(0.35),
        [("理论支撑：维果茨基最近发展区  ·  支架式教学", 13, False, OCHRE)],
        PP_ALIGN.CENTER,
    )


# ── Page 4: 评价·多元化育人 ──
def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    hdr(s, "评价 · 多元化育人", "明镜照心：在教-学-评一体中见证生命生长", 4)
    tb(
        s,
        PL,
        Inches(1.0),
        IW,
        Inches(0.45),
        [("评价哲学：告别终场裁判，让评价成为贯穿全流程的探照灯", 14, True, DAILAN)],
        PP_ALIGN.CENTER,
    )
    dims = [
        ("课前", "唤醒经验", "情境激趣"),
        ("课中", "嵌入任务", "人物证据卡"),
        ("课终", "表现创作", "「想象说明卡」+ 星级量表"),
    ]
    cw = (IW - Inches(0.3)) / 3
    x = PL
    for phase, action, detail in dims:
        verse_card(
            s,
            x,
            Inches(1.65),
            cw,
            Inches(2.35),
            [
                (phase, 16, True, RED),
                (action, 14, True, DAILAN),
                ("", 6, False, INK),
                (detail, 13, False, INK),
            ],
            DAILAN,
            DAILAN_WASH,
        )
        x += cw + Inches(0.15)
    verse_card(
        s,
        PL,
        Inches(4.25),
        IW,
        Inches(2.05),
        [
            ("多元互评机制", 14, True, BAMBOO),
            ("依据「依据、逻辑、立意、文采」四维量表自评互评", 13, False, INK),
            ("", 6, False, INK),
            ("促使学生从「被评价者」转变为「自我矫正的思考者」", 14, True, RED),
        ],
        OCHRE,
        BAMBOO_WASH,
    )


# ── Page 5: 结语·育人本色 ──
def s05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    hdr(s, "结语 · 育人本色", "文以化人：让语文素养在真实与创造中扎根", 5)
    wash(s, Inches(1.5), Inches(1.05), Inches(10.3), Inches(3.5), DAILAN_WASH)
    verse_card(
        s,
        PL,
        Inches(1.15),
        IW,
        Inches(2.55),
        [
            ("核心反思", 14, True, RED),
            ("", 6, False, INK),
            ("读奇幻故事，并非逃避现实，而是为了洞察生活的真相；", 15, False, INK),
            ("经历语文实践，不只积累语言，更为培养独立思考与纯真批判。", 15, False, INK),
        ],
        DAILAN,
        BG,
    )
    cloud_div(s, Inches(4.05))
    line_box(s, PL + Inches(0.8), Inches(4.35), IW - Inches(1.6), Inches(1.85), BG, RED, Pt(1.5))
    tb(
        s,
        PL + Inches(1.0),
        Inches(4.55),
        IW - Inches(2.0),
        Inches(1.5),
        [
            ("卷尾落款", 12, True, OCHRE),
            ("", 8, False, INK),
            ("「让一个想象从书页中继续生长，", 17, True, DAILAN),
            ("也让一个生命在语言的润泽下巍然拔节。」", 17, True, DAILAN),
        ],
        PP_ALIGN.CENTER,
    )
    tb(
        s,
        PL,
        Inches(6.55),
        IW,
        Inches(0.35),
        [("想象与真实 · 单元学历案 · 教研汇报", 12, False, MUTED)],
        PP_ALIGN.CENTER,
    )


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    for fn in (s01, s02, s03, s04, s05):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    import shutil

    shutil.copy2(OUT, OUT_EN)
    print(f"✓ {OUT}")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    main()
