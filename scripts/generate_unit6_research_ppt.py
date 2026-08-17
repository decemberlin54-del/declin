#!/usr/bin/env python3
"""教研汇报 PPT：想象与真实——单元学历案设计与实践（6 slides）."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/单元学历案教研汇报.pptx")
OUT_EN = Path("/workspace/unit6-lesson-plan-research-report.pptx")

SW = Inches(13.333)
SH = Inches(7.5)
PL, PT = Inches(0.55), Inches(0.4)
IW = SW - Inches(1.1)

# 现代商务/教育配色
NAVY = RGBColor(0x1A, 0x3A, 0x5C)       # 深蓝
GREEN = RGBColor(0x2D, 0x5A, 0x4A)      # 墨绿
GREEN_SOFT = RGBColor(0xE8, 0xF0, 0xEC)
NAVY_SOFT = RGBColor(0xE9, 0xEE, 0xF4)
ORANGE = RGBColor(0xE0, 0x7B, 0x39)      # 暖橙
YELLOW = RGBColor(0xFF, 0xC1, 0x07)       # 明黄
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x6B, 0x7B, 0x8C)


def font(run, size=16, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide, color=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    tree = slide.shapes._spTree
    tree.remove(sh._element)
    tree.insert(2, sh._element)


def accent_bar(slide, color=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), SH)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def card(slide, l, t, w, h, fill=WHITE, border=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(1.2)
    return sh


def tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(item, tuple):
            txt, sz, b, c = item
            r = p.add_run()
            font(r, sz, b, c)
            r.text = txt
        else:
            r = p.add_run()
            font(r)
            r.text = item
        p.space_after = Pt(4)
    return box


def slide_num(slide, n, total=6):
    tb(slide, SW - Inches(1.0), SH - Inches(0.45), Inches(0.6), Inches(0.3), [(f"{n}/{total}", 11, False, MUTED)], PP_ALIGN.RIGHT)


def s01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    accent_bar(s, ORANGE)
    # 装饰块
    card(s, Inches(0.5), Inches(5.8), Inches(4.5), Inches(0.08), ORANGE, ORANGE)
    tb(
        s,
        Inches(0.75),
        Inches(1.4),
        Inches(11.5),
        Inches(1.4),
        [("想象与真实：基于核心素养的单元学历案设计与实践", 34, True, WHITE)],
    )
    tb(
        s,
        Inches(0.75),
        Inches(2.85),
        Inches(11.0),
        Inches(0.55),
        [("统编版语文七年级上册第六单元教学设计汇报", 18, False, RGBColor(0xB8, 0xD4, 0xE8))],
    )
    tb(s, Inches(0.75), Inches(3.55), Inches(4.0), Inches(0.45), [("汇报人：朱琳", 16, False, WHITE)])
    card(s, Inches(0.75), Inches(4.35), Inches(9.5), Inches(0.65), GREEN, GREEN)
    tb(
        s,
        Inches(0.95),
        Inches(4.45),
        Inches(9.1),
        Inches(0.45),
        [("设计理念：理论支撑 · 情境驱动 · 支架赋能 · 评教一体", 15, True, WHITE)],
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )


def s02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    accent_bar(s, GREEN)
    tb(s, PL + Inches(0.15), PT, IW, Inches(0.55), [("理论基石与设计逻辑", 26, True, NAVY)])
    tb(s, PL + Inches(0.15), PT + Inches(0.5), IW, Inches(0.35), [("Theoretical Foundations", 12, False, MUTED)])
    theories = [
        ("逆向教学设计 (UbD)", "威金斯 & 麦克泰格", "以终为始：目标 → 评估证据 → 学习体验", NAVY, NAVY_SOFT),
        ("情境学习理论", "让·莱夫 & 埃蒂安·温格", "真实角色驱动：「想象力博物馆」策展人", GREEN, GREEN_SOFT),
        ("支架式教学理论", "维果茨基最近发展区", "显性思维工具：推理图、证据卡、联想桥", NAVY, NAVY_SOFT),
        ("《语文课程标准》", "实践导向", "核心素养统领，学习任务群集成整合", GREEN, GREEN_SOFT),
    ]
    cw = (IW - Inches(0.15)) / 2
    ch = Inches(2.35)
    for i, (title, author, desc, border, fill) in enumerate(theories):
        col, row = i % 2, i // 2
        x = PL + Inches(0.15) + col * (cw + Inches(0.15))
        y = Inches(1.15) + row * (ch + Inches(0.15))
        card(s, x, y, cw, ch, fill, border)
        tb(
            s,
            x + Inches(0.18),
            y + Inches(0.15),
            cw - Inches(0.3),
            ch - Inches(0.25),
            [
                (title, 15, True, border),
                (author, 11, False, ORANGE),
                ("", 4, False, TEXT),
                (desc, 13, False, TEXT),
            ],
        )
    slide_num(s, 2)


def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    accent_bar(s, NAVY)
    tb(s, PL + Inches(0.15), PT, IW, Inches(0.55), [("维度一：创新课程设计——大观念与大任务", 24, True, NAVY)])
    card(s, PL + Inches(0.15), Inches(1.05), IW, Inches(1.05), NAVY_SOFT, NAVY)
    tb(
        s,
        PL + Inches(0.35),
        Inches(1.2),
        IW - Inches(0.5),
        Inches(0.85),
        [
            ("核心大观念", 13, True, ORANGE),
            ("整体重构单元，探究「想象与真实」的隐秘联系", 16, True, TEXT),
            ("大任务情境：「想象力博物馆」策展人（9课时整体规划）", 14, False, MUTED),
        ],
    )
    cw = (IW - Inches(0.15)) / 2
    card(s, PL + Inches(0.15), Inches(2.3), cw, Inches(4.35), GREEN_SOFT, GREEN)
    tb(
        s,
        PL + Inches(0.35),
        Inches(2.45),
        cw - Inches(0.35),
        Inches(4.0),
        [
            ("任务一 · 走进想象世界（1—7课时）", 14, True, GREEN),
            ("", 4, False, TEXT),
            ("《小圣施威降大圣》—— 自由与规则", 13, False, TEXT),
            ("《皇帝的新装》—— 真实与虚假", 13, False, TEXT),
            ("《女娲造人》—— 创造与生命", 13, False, TEXT),
            ("《寓言四则》—— 智慧与局限", 13, False, TEXT),
        ],
    )
    card(s, PL + cw + Inches(0.3), Inches(2.3), cw, Inches(4.35), NAVY_SOFT, ORANGE)
    tb(
        s,
        PL + cw + Inches(0.5),
        Inches(2.45),
        cw - Inches(0.35),
        Inches(4.0),
        [
            ("任务二 · 把想象变成作品（8—9课时）", 14, True, ORANGE),
            ("", 4, False, TEXT),
            ("制作「想象说明卡」", 15, True, TEXT),
            ("", 6, False, TEXT),
            ("创意写作 / 续写 / 改写 / 课本剧实践", 13, False, TEXT),
            ("", 8, False, TEXT),
            ("双重任务递进：读懂 → 创编", 13, True, GREEN),
        ],
    )
    slide_num(s, 3)


def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    accent_bar(s, GREEN)
    tb(s, PL + Inches(0.15), PT, IW, Inches(0.55), [("维度二：以学生为主体——显性认知支架", 24, True, NAVY)])
    card(s, PL + Inches(0.15), Inches(1.0), IW, Inches(0.75), YELLOW, ORANGE)
    tb(
        s,
        PL + Inches(0.35),
        Inches(1.12),
        IW - Inches(0.5),
        Inches(0.55),
        [("问题痛点：七年级学生易沉溺于「情节热闹」，忽视「隐喻真实」", 14, True, TEXT)],
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )
    scaffolds = [
        ("01", "阅读速度与概括支架", "「节点词+简洁句」快速默读表（≥400字/分）"),
        ("02", "深度思维推理支架", "七十二变「环境—形态—对手」限度推理图"),
        ("03", "文本鉴赏分析支架", "人物鉴赏「证据—特点—作用」三步表达法"),
        ("04", "创意表达迁移支架", "写作发散「联想桥」四维模型（形状/功能/情感/场景）"),
    ]
    y = Inches(2.0)
    for num, title, desc in scaffolds:
        card(s, PL + Inches(0.15), y, Inches(0.65), Inches(0.95), NAVY, NAVY)
        tb(s, PL + Inches(0.15), y + Inches(0.22), Inches(0.65), Inches(0.5), [(num, 14, True, WHITE)], PP_ALIGN.CENTER)
        card(s, PL + Inches(0.95), y, IW - Inches(1.1), Inches(0.95), WHITE, GREEN if num in ("02", "04") else NAVY)
        tb(
            s,
            PL + Inches(1.1),
            y + Inches(0.1),
            IW - Inches(1.35),
            Inches(0.75),
            [(title, 14, True, NAVY), (desc, 12, False, MUTED)],
        )
        y += Inches(1.1)
    slide_num(s, 4)


def s05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    accent_bar(s, NAVY)
    tb(s, PL + Inches(0.15), PT, IW, Inches(0.55), [("维度三：评价多元化——教-学-评一体化", 24, True, NAVY)])
    evals = [
        ("诊断性评价", "课前热身情境思考，激活既有经验", NAVY, NAVY_SOFT),
        ("形成性评价", "课堂嵌入式「评价任务」+ 人物证据卡", GREEN, GREEN_SOFT),
        ("总结性评价", "终极成果「想象说明卡」+ 创作作品量表", ORANGE, RGBColor(0xFD, 0xF0, 0xE6)),
    ]
    cw = (IW - Inches(0.3)) / 3
    x = PL + Inches(0.15)
    for title, desc, border, fill in evals:
        card(s, x, Inches(1.05), cw, Inches(2.5), fill, border)
        tb(
            s,
            x + Inches(0.15),
            Inches(1.2),
            cw - Inches(0.25),
            Inches(2.2),
            [(title, 15, True, border), ("", 6, False, TEXT), (desc, 13, False, TEXT)],
            PP_ALIGN.CENTER,
        )
        x += cw + Inches(0.15)
    card(s, PL + Inches(0.15), Inches(3.85), IW, Inches(2.55), GREEN_SOFT, GREEN)
    tb(
        s,
        PL + Inches(0.35),
        Inches(4.0),
        IW - Inches(0.5),
        Inches(2.25),
        [
            ("多元主体参与", 14, True, GREEN),
            ("制定细化的自评与互评星级量表", 13, False, TEXT),
            ("文本依据 / 想象逻辑 / 主题提升 / 语言表达", 13, True, ORANGE),
            ("", 6, False, TEXT),
            ("实现「以评促学」「以评促改」", 15, True, NAVY),
        ],
        PP_ALIGN.CENTER,
    )
    slide_num(s, 5)


def s06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    accent_bar(s, ORANGE)
    tb(s, Inches(0.75), Inches(0.55), Inches(11.0), Inches(0.55), [("总结与反思", 26, True, WHITE)])
    tb(s, Inches(0.75), Inches(1.0), Inches(11.0), Inches(0.35), [("Conclusion", 12, False, RGBColor(0xB8, 0xD4, 0xE8))])
    shifts = [
        ("教学方式", "从「教课文」", "走向「引导语文实践活动」"),
        ("评价方式", "从「单向测试」", "走向「多元表现性评价」"),
        ("课堂主体", "从「教师主导」", "走向「学生自主建构」"),
    ]
    y = Inches(1.55)
    for label, before, after in shifts:
        card(s, Inches(0.75), y, Inches(11.8), Inches(1.05), RGBColor(0x22, 0x4A, 0x6E), GREEN)
        tb(
            s,
            Inches(0.95),
            y + Inches(0.15),
            Inches(11.4),
            Inches(0.75),
            [
                (label, 13, True, YELLOW),
                (before, 14, False, RGBColor(0xB8, 0xD4, 0xE8)),
                ("  →  ", 14, True, ORANGE),
                (after, 14, True, WHITE),
            ],
        )
        y += Inches(1.2)
    card(s, Inches(0.75), Inches(5.35), Inches(11.8), Inches(1.35), GREEN, ORANGE)
    tb(
        s,
        Inches(0.95),
        Inches(5.55),
        Inches(11.4),
        Inches(1.0),
        [
            ("核心落脚点", 13, True, YELLOW),
            ("在文学想象的光影中照见生活，在真实的语文实践中", 15, True, WHITE),
            ("培养独立批判与创造能力。", 15, True, WHITE),
        ],
        PP_ALIGN.CENTER,
    )
    slide_num(s, 6)


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    for fn in (s01, s02, s03, s04, s05, s06):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    import shutil

    shutil.copy2(OUT, OUT_EN)
    print(f"✓ {OUT}")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    main()
