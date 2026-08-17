#!/usr/bin/env python3
"""第24课《寓言四则》国风水墨风格 + 渐显动画 PPT — 9 slides."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from ppt_anim import Build, slide_transition  # noqa: E402

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/第24课 寓言四则.pptx")
OUT_EN = Path("/workspace/unit6-lesson24-fables.pptx")
OUT2 = Path("/workspace/第六单元PPT/寓言四则短故事大镜子.pptx")

SW = Inches(13.333)
SH = Inches(7.5)
PL, PT = Inches(0.55), Inches(0.35)
IW = SW - Inches(1.1)

BG = RGBColor(0xF4, 0xF1, 0xEA)
GREEN = RGBColor(0x4A, 0x7C, 0x6F)
GREEN_SOFT = RGBColor(0xE4, 0xED, 0xE8)
WASH = RGBColor(0xD8, 0xE6, 0xDF)
INK = RGBColor(0x1E, 0x1E, 0x1E)
RED = RGBColor(0xC8, 0x3C, 0x23)
MUTED = RGBColor(0x5A, 0x6A, 0x62)
FT = "KaiTi"
FB = "Microsoft YaHei"

ARCHIVE = [
    ("赫耳墨斯和雕像者", "赫耳墨斯、雕像者", "问价自取其辱", "讽刺妄自尊大"),
    ("蚊子和狮子", "蚊子、狮子", "胜后被蛛网困", "骄兵必败"),
    ("穿井得一人", "丁氏、宋君", "传言失真求证", "以讹传讹须求证"),
    ("杞人忧天", "杞人、晓之者", "忧天被开导", "不必杞人忧天"),
]


def fnt(run, size=16, bold=False, color=INK, title=False):
    run.font.name = FT if title else FB
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    sh.fill.solid()
    sh.fill.fore_color.rgb = BG
    sh.line.fill.background()
    t = slide.shapes._spTree
    t.remove(sh._element)
    t.insert(2, sh._element)


def wash(slide, l, t, w, h, c=WASH):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = c
    sh.line.fill.background()


def lbox(slide, l, t, w, h, fill=BG, border=GREEN):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border
        sh.line.width = Pt(1.5)
    return sh


def bar(slide, l, t, h, c=GREEN, w=Inches(0.05)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = c
    sh.line.fill.background()


def seal(slide, l, t, sz, txt, fs=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, sz, sz)
    sh.fill.background()
    sh.line.color.rgb = RED
    sh.line.width = Pt(2)
    tb(slide, l, t, sz, sz, [(txt, fs, True, RED)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def cloud(slide, y):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PL, y, IW, Inches(0.01))
    ln.fill.solid()
    ln.fill.fore_color.rgb = MUTED
    ln.line.fill.background()
    tb(slide, PL, y - Inches(0.06), IW, Inches(0.2), [("～ ～ ～ ～ ～ ～ ～ ～ ～ ～", 9, False, MUTED)], PP_ALIGN.CENTER)


def tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, title=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(item, tuple):
            txt, sz, b, c = item
            r = p.add_run()
            fnt(r, sz, b, c, title=title or b)
            r.text = txt
        else:
            r = p.add_run()
            fnt(r, title=title)
            r.text = item
        p.space_after = Pt(3)
    return box


def hdr(slide, title, n):
    seal(slide, PL, PT, Inches(0.38), f"{n:02d}", 10)
    tb(slide, PL + Inches(0.48), PT, IW * 0.7, Inches(0.38), [(title, 19, True, INK)], title=True)
    cloud(slide, PT + Inches(0.44))


def card(slide, l, t, w, h, lines, accent=True, fill=BG):
    if fill != BG:
        wash(slide, l, t, w, h, fill)
    else:
        lbox(slide, l, t, w, h, BG, GREEN)
    if accent:
        bar(slide, l, t, h)
    tb(slide, l + Inches(0.14), t + Inches(0.1), w - Inches(0.2), h - Inches(0.15), lines)


def s01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.4)
    wash(s, Inches(0.2), Inches(0.3), Inches(7.0), Inches(6.8), WASH)
    wash(s, Inches(8.5), Inches(1.0), Inches(4.5), Inches(5.5), GREEN_SOFT)
    # 主标题 static
    tb(s, PL, Inches(1.2), Inches(7.5), Inches(2.0), [("寓言四则", 56, True, INK)], title=True)
    b = Build(s)
    with b.step(auto=True, delay_ms=300):
        tb(
            s,
            Inches(9.0),
            Inches(1.5),
            Inches(3.5),
            Inches(4.5),
            [
                ("短", 32, True, INK),
                ("故", 32, True, INK),
                ("事", 32, True, INK),
                ("", 8, False, INK),
                ("大", 32, True, RED),
                ("镜", 32, True, RED),
                ("子", 32, True, RED),
                ("", 10, False, INK),
                ("照见人心", 16, False, MUTED),
                ("折射现实", 16, False, MUTED),
            ],
            PP_ALIGN.CENTER,
        )
        seal(s, Inches(11.8), Inches(5.8), Inches(0.45), "寓言", 10)
    b.apply(effect="fade", duration=0.6)
    tb(s, PL, Inches(6.5), IW, Inches(0.35), [("课时六、七 · 短故事，大镜子", 13, False, GREEN)], PP_ALIGN.CENTER)


def s02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "情境导入与课时目标", 2)
    cw = (IW - Inches(0.2)) / 2
    b = Build(s)
    with b.step():
        card(
            s,
            PL,
            Inches(0.62),
            cw,
            Inches(1.55),
            [
                ("追问一", 12, True, RED),
                ("为什么《赫耳墨斯》中雕像者的一句话会让他难堪？", 15, True, INK),
            ],
        )
    with b.step():
        card(
            s,
            PL,
            Inches(2.3),
            cw,
            Inches(1.55),
            [
                ("追问二", 12, True, RED),
                ("为什么《穿井得一人》的一句话会在国都传开？", 15, True, INK),
            ],
            fill=GREEN_SOFT,
        )
    with b.step():
        card(
            s,
            PL + cw + Inches(0.2),
            Inches(0.62),
            cw,
            Inches(5.5),
            [
                ("课时目标", 15, True, GREEN),
                ("1. 归纳寓言文体特点，理解各则寓意", 14, False, INK),
                ("2. 疏通文言，积累重点词语", 14, False, INK),
                ("3. 分析情节与寓意关系，探究多种理解", 14, False, INK),
                ("", 8, False, INK),
                ("评价任务", 15, True, RED),
                ("① 完成寓言档案卡", 14, False, INK),
                ("② 翻译重点句，解释文言词", 14, False, INK),
                ("③ 分析寓意，联系生活谈启示", 14, False, INK),
            ],
        )
    b.apply(effects=["fade", "fade", "wipe_down"], duration=0.5)


def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "寓言档案卡", 3)
    cw = (IW - Inches(0.25)) / 4
    b = Build(s)
    for i, (name, people, plot, moral) in enumerate(ARCHIVE):
        x = PL + i * (cw + Inches(0.08))
        with b.step():
            fill = GREEN_SOFT if i % 2 else BG
            card(
                s,
                x,
                Inches(0.65),
                cw,
                Inches(5.7),
                [
                    (name, 13, True, GREEN),
                    ("", 4, False, INK),
                    ("人物", 11, True, RED),
                    (people, 12, False, INK),
                    ("情节", 11, True, RED),
                    (plot, 12, False, INK),
                    ("寓意", 11, True, RED),
                    (moral, 12, True, INK),
                ],
                fill=fill,
            )
    b.apply(effect="wipe_right", duration=0.45)


def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "寓言文体特点", 4)
    items = [
        ("篇幅", "短小精悍"),
        ("人物", "典型鲜明"),
        ("写法", "拟人 · 夸张 · 反转"),
        ("作用", "寄寓道理 · 警醒世人"),
    ]
    b = Build(s)
    w = Inches(2.85)
    with b.step():
        for i, (k, v) in enumerate(items[:3]):
            x = PL + i * (w + Inches(0.15))
            wash(s, x, Inches(0.75), w, Inches(2.2), BG if i % 2 else GREEN_SOFT)
            bar(s, x, Inches(0.75), Inches(2.2))
            tb(s, x + Inches(0.12), Inches(0.95), w - Inches(0.2), Inches(1.8), [(k, 16, True, GREEN), (v, 15, False, INK)])
            if i < 2:
                tb(s, x + w, Inches(1.5), Inches(0.2), Inches(0.4), [("→", 18, True, GREEN)])
    with b.step():
        x = PL + 3 * (w + Inches(0.15))
        lbox(s, x, Inches(0.75), w, Inches(2.2), RED, RED)
        tb(s, x + Inches(0.12), Inches(0.95), w - Inches(0.2), Inches(1.8), [("作用", 16, True, RED), ("寄寓道理 · 警醒世人", 15, True, INK)])
    b.apply(effects=["wipe_right", "fade"], duration=0.5)


def s05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "读懂寓意与改写探究", 5)
    b = Build(s)
    with b.step():
        card(
            s,
            PL,
            Inches(0.65),
            IW,
            Inches(2.55),
            [
                ("情节 → 寓意", 16, True, GREEN),
                ("赫耳墨斯：「白送」→ 讽刺自高自大", 14, False, INK),
                ("蚊子：凯歌后被粘 → 骄兵必败", 14, False, INK),
                ("丁氏：澄清误会 → 勿以讹传讹", 14, False, INK),
            ],
            fill=GREEN_SOFT,
        )
    with b.step():
        lbox(s, PL, Inches(3.45), IW, Inches(2.85), RED, RED)
        bar(s, PL, Inches(3.45), Inches(2.85), RED, Inches(0.06))
        tb(
            s,
            PL + Inches(0.2),
            Inches(3.65),
            IW - Inches(0.35),
            Inches(2.5),
            [
                ("改写探究", 15, True, RED),
                ("若蚊子战胜后悄悄离开……", 16, True, INK),
                ("寓意或变为「懂得适可而止」", 15, False, INK),
                ("说明：寓意与情节设计密切相关", 14, True, GREEN),
            ],
        )
    b.apply(effects=["fade", "fly_up"], duration=0.55)


def s06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "联系生活", 6)
    scenes = [
        ("①", "未经核实转发消息", "《穿井得一人》"),
        ("②", "成绩好看不起同学", "《赫耳墨斯》"),
        ("③", "听到风声就焦虑", "《杞人忧天》"),
    ]
    cw = (IW - Inches(0.25)) / 3
    b = Build(s)
    for i, (num, scene, fable) in enumerate(scenes):
        with b.step():
            x = PL + i * (cw + Inches(0.12))
            card(
                s,
                x,
                Inches(0.65),
                cw,
                Inches(5.5),
                [
                    (num, 28, True, RED),
                    ("", 6, False, INK),
                    (scene, 16, True, INK),
                    ("", 10, False, INK),
                    ("对位寓言", 12, True, GREEN),
                    (fable, 18, True, RED),
                ],
                fill=GREEN_SOFT if i == 1 else BG,
            )
    b.apply(effect="fade", duration=0.5)


def s07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "通文言，积词语", 7)
    words = [("闻", "听说"), ("道", "讲述"), ("亡", "同「无」"), ("晓", "开导"), ("舍然", "释然")]
    trans = [
        "（1）得一人之使，非得一人于井中也。",
        "（2）求闻之若此，不若无闻也。",
        "（3）若屈伸呼吸，终日在天中行止，奈何忧崩坠乎？",
    ]
    b = Build(s)
    with b.step():
        y = Inches(0.65)
        for w, m in words:
            lbox(s, PL, y, Inches(4.8), Inches(0.72), BG, GREEN)
            bar(s, PL, y, Inches(0.72))
            tb(s, PL + Inches(0.15), y + Inches(0.1), Inches(4.5), Inches(0.55), [(f"{w}：{m}", 15, True, INK)])
            y += Inches(0.82)
    with b.step():
        y = Inches(0.65)
        for t in trans:
            lbox(s, PL + Inches(5.2), y, Inches(6.8), Inches(1.35), GREEN_SOFT, GREEN)
            tb(s, PL + Inches(5.35), y + Inches(0.15), Inches(6.5), Inches(1.05), [(t, 14, False, INK)])
            y += Inches(1.45)
    b.apply(effects=["fade", "wipe_right"], duration=0.5)


def s08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "「杞人忧天」的多视角解读", 8)
    cw = (IW - Inches(0.2)) / 2
    b = Build(s)
    with b.step():
        card(
            s,
            PL,
            Inches(0.65),
            cw,
            Inches(5.6),
            [
                ("视角 A · 讽刺不必要的杞忧", 15, True, GREEN),
                ("依据：「废寝食者」「奈何忧崩坠乎」", 14, False, INK),
                ("启示：遇事应实事求是，勿凭空忧虑", 14, True, INK),
            ],
            fill=GREEN_SOFT,
        )
    with b.step():
        card(
            s,
            PL + cw + Inches(0.2),
            Inches(0.65),
            cw,
            Inches(5.6),
            [
                ("视角 B · 体现忧患意识", 15, True, RED),
                ("依据：「忧天地崩坠，身亡所寄」", 14, False, INK),
                ("启示：对未知保持审慎，亦需理性求证", 14, True, INK),
            ],
        )
    b.apply(effects=["fly_left", "fly_right"], duration=0.55)


def s09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    slide_transition(s, "fade", 0.35)
    hdr(s, "给古人一条今天的建议", 9)
    steps = ["丁氏原话", "传言失真", "国人传播", "宋君求证"]
    b = Build(s)
    with b.step():
        x = PL
        sw = Inches(2.55)
        for i, st in enumerate(steps):
            lbox(s, x, Inches(0.75), sw, Inches(0.85), BG, GREEN)
            tb(s, x, Inches(0.88), sw, Inches(0.6), [(st, 13, True, INK)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            if i < 3:
                tb(s, x + sw, Inches(0.95), Inches(0.35), Inches(0.4), [("→", 16, True, RED)], PP_ALIGN.CENTER)
            x += sw + Inches(0.35)
        tb(
            s,
            PL,
            Inches(1.85),
            Inches(6.5),
            Inches(0.55),
            [("面对网络信息：核对来源 · 理性判断 · 不传谣言", 14, True, GREEN)],
        )
    with b.step():
        lbox(s, PL + Inches(6.8), Inches(0.65), Inches(5.5), Inches(5.6), RED, RED)
        bar(s, PL + Inches(6.8), Inches(0.65), Inches(5.6), RED, Inches(0.06))
        tb(
            s,
            PL + Inches(7.0),
            Inches(1.0),
            Inches(5.1),
            Inches(4.8),
            [
                ("课堂小结", 18, True, RED),
                ("", 8, False, INK),
                ("信息须求证", 22, True, INK),
                ("忧虑须理性", 22, True, INK),
                ("智慧照见局限", 22, True, INK),
            ],
            PP_ALIGN.CENTER,
        )
    b.apply(effects=["wipe_right", "fade"], duration=0.55)


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    import shutil

    shutil.copy2(OUT, OUT_EN)
    shutil.copy2(OUT, OUT2)
    print(f"✓ {OUT}")
    print(f"✓ {OUT_EN}")
    print(f"✓ {OUT2}")


if __name__ == "__main__":
    main()
