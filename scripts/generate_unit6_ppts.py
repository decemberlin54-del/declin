#!/usr/bin/env python3
"""Generate Unit 6 lesson PPTs aligned with 学历案."""

import os
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path("/workspace/第六单元PPT")
IMG_DIR = BASE / "images"
OUT_DIR = BASE

# Morandi + Dopamine palette
M = {
    "rose": (196, 164, 164),
    "sage": (163, 177, 165),
    "sky": (164, 180, 196),
    "lav": (180, 170, 196),
    "sand": (210, 198, 180),
    "slate": (90, 96, 110),
    "white": (252, 250, 247),
    "coral": (255, 127, 102),
    "mint": (102, 205, 170),
    "sun": (255, 209, 102),
    "peach": (255, 183, 153),
    "lilac": (200, 162, 220),
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def ensure_bg(name: str, accent: tuple, label: str) -> Path:
    path = IMG_DIR / name
    if path.exists():
        return path
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1920, 1080), M["white"])
    draw = ImageDraw.Draw(img)
    for i in range(0, 1080, 4):
        t = i / 1080
        c = tuple(int(M["white"][j] * (1 - t * 0.15) + accent[j] * t * 0.15) for j in range(3))
        draw.line([(0, i), (1920, i)], fill=c)
    draw.ellipse((1400, -120, 2100, 580), fill=(*accent, 40) if len(accent) == 3 else accent)
    draw.ellipse((-200, 700, 500, 1200), fill=(*M["sand"],))
    img.save(path)
    return path


def rgb(t):
    return RGBColor(*t)


def set_run(run, text, size=28, bold=False, color=M["slate"]):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.color.rgb = rgb(color)


def add_bg(slide, bg_path: Path, alpha=0.18):
    slide.shapes.add_picture(str(bg_path), 0, 0, SLIDE_W, SLIDE_H)
    shape = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(M["white"])
    shape.fill.transparency = 1 - alpha
    shape.line.fill.background()


def add_bar(slide, color, top=Inches(0), height=Inches(0.12)):
    bar = slide.shapes.add_shape(1, 0, top, SLIDE_W, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(color)
    bar.line.fill.background()


def add_title_block(slide, title, subtitle=None, accent=M["coral"]):
    add_bar(slide, accent)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.8), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    set_run(p.add_run(), title, 44, True, M["slate"])
    if subtitle:
        p2 = tf.add_paragraph()
        set_run(p2.add_run(), subtitle, 26, False, accent)


def add_content_slide(prs, bg, title, bullets, accent, tag="学习任务"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_bar(slide, accent, Inches(0), Inches(0.1))
    tag_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.25), Inches(3), Inches(0.5))
    set_run(tag_box.text_frame.paragraphs[0].add_run(), tag, 20, True, accent)
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.75), Inches(11.5), Inches(1.0))
    set_run(title_box.text_frame.paragraphs[0].add_run(), title, 36, True, M["slate"])
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.85), Inches(11.2), Inches(4.8))
    tf = body.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        p.level = 0
        if isinstance(item, tuple):
            label, text = item
            set_run(p.add_run(), f"{label}  ", 26, True, accent)
            set_run(p.add_run(), text, 26, False, M["slate"])
        else:
            set_run(p.add_run(), item, 28, False, M["slate"])
    return slide


def add_question_slide(prs, bg, question, hint, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_bar(slide, accent)
    qbox = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(2.5))
    tf = qbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), "💭 想一想", 30, True, accent)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)
    set_run(p2.add_run(), question, 34, True, M["slate"])
    if hint:
        hbox = slide.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10.0), Inches(1.2))
        hp = hbox.text_frame.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        set_run(hp.add_run(), hint, 24, False, accent)
    return slide


def add_cover(prs, bg, main, sub, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg, 0.25)
    add_bar(slide, accent, Inches(3.2), Inches(0.08))
    mbox = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.0), Inches(1.5))
    mp = mbox.text_frame.paragraphs[0]
    mp.alignment = PP_ALIGN.CENTER
    set_run(mp.add_run(), main, 52, True, M["slate"])
    sbox = slide.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.0), Inches(1.0))
    sp = sbox.text_frame.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    set_run(sp.add_run(), sub, 28, False, accent)


def add_summary_slide(prs, bg, lines, accent):
    add_content_slide(
        prs, bg, "课堂小结", lines, accent, tag="回顾收获"
    )


def save_ppt(prs, filename):
  OUT_DIR.mkdir(parents=True, exist_ok=True)
  path = OUT_DIR / filename
  prs.save(str(path))
  print(f"✓ {path} ({len(prs.slides)} slides)")


def ppt_unit_intro():
    bg = ensure_bg("bg_unit.png", M["sky"], "unit")
    accent = M["sky"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "想象力博物馆", "七年级上册 · 第六单元 · 第一课时", accent)

    add_content_slide(prs, bg, "本课目标", [
        ("目标1", "感知单元学习目标、文体类型与挑战任务"),
        ("目标2", "快速默读，用关键词概括情节"),
    ], accent, "课时目标")

    add_content_slide(prs, bg, "大任务：想象力博物馆", [
        "校园文学社微展览，你是策展人",
        "制作「想象说明卡」",
        "完成原创想象文字或微型课本剧",
    ], M["coral"], "单元挑战")

    add_content_slide(prs, bg, "母题递进链", [
        ("自由与规则", "《小圣施威降大圣》"),
        ("真实与虚假", "《皇帝的新装》"),
        ("创造与生命", "《女娲造人》"),
        ("智慧与局限", "《寓言四则》"),
    ], M["lav"], "想象与真实")

    add_question_slide(
        prs, bg,
        "如果伞能说话、星星掉进教室……\n这些离奇念头从哪来？",
        "提示：从真实经验出发，联想与想象",
        M["mint"],
    )

    add_content_slide(prs, bg, "学习任务一 · 明确任务", [
        "阅读学历案，圈画挑战任务与学习目标",
        "同伴交流：说说你的策展想法",
        "填写：本单元文体？阅读方法？",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "学习任务二 · 读得快也读得准", [
        "快速默读（不出声、不指读）",
        "段末停顿，用5—8字记录段意",
        "归纳复述三要素：人物、＿＿、＿＿",
    ], M["sun"], "快速阅读")

    add_question_slide(
        prs, bg,
        "同桌互评：你说清\n人物、事件、结果了吗？",
        "每分钟阅读不少于400字",
        accent,
    )

    add_summary_slide(prs, bg, [
        "本单元学习重点：想象与真实",
        "四条策展主题：自由→真实→创造→智慧",
        "快速阅读：人物—事件—结果",
    ], accent)

    add_content_slide(prs, bg, "作业", [
        ("必做", "预习一篇课文，写100字以内梗概"),
        ("选做", "记录一个「如果……会怎样」的想象问题"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第1课时 单元导学.pptx")


def ppt_lesson21():
    bg = ensure_bg("bg_wukong.png", M["coral"], "wukong")
    accent = M["coral"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "小圣施威降大圣", "第21课 · 第二、三课时 · 自由与规则", accent)

    add_content_slide(prs, bg, "本课目标", [
        ("目标1", "梳理孙悟空与二郎神的变化链"),
        ("目标2", "分析神魔形象与想象趣味"),
        ("目标3", "探究「自由与规则」母题"),
    ], accent, "课时目标")

    add_question_slide(
        prs, bg,
        "高手过招，吸引人的不只是\n「谁赢了」——",
        "把自己当作赛事解说员，抓住每一次变化",
        M["sun"],
    )

    add_content_slide(prs, bg, "学习任务一 · 变形追逐图", [
        "默读课文，圈出每次「变作」对象",
        "填写七十二变对照表",
        "用句式复述最精彩一轮（≤40字）",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "学习任务二 · 加解说词", [
        "变庙宇：尾巴为何变旗竿？",
        "完成人物档案卡（证据—特点—作用）",
        "批注鲁迅、林庚的评论",
    ], M["mint"], "细读品味")

    add_question_slide(
        prs, bg,
        "「使神魔皆有人情」\n课文哪里体现了人情与世故？",
        "林庚：不宜看得过于认真，应看到儿童心理",
        accent,
    )

    add_content_slide(prs, bg, "学习任务三 · 提炼策展主题", [
        "讨论：孙悟空为何逃？天庭为何捉？",
        "联系生活：自由与规则矛盾吗？",
        "主题一：自由可贵，但＿＿＿＿",
    ], M["lav"], "主题探究")

    add_summary_slide(prs, bg, [
        "想象精彩，来自变化中的因果回应",
        "神魔有法力，也有人情味",
        "策展主题：自由与规则",
    ], accent)

    add_content_slide(prs, bg, "作业", [
        ("必做", "写80字「赛事快讯」（≥3次变化）"),
        ("链接", "作者吴承恩，成书明代，再写一个经典情节"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第21课 小圣施威降大圣.pptx")


def ppt_lesson22():
    bg = ensure_bg("bg_emperor.png", M["lav"], "emperor")
    accent = M["lav"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "皇帝的新装", "第22课 · 第四、五课时 · 真实与虚假", accent)

    add_content_slide(prs, bg, "本课目标", [
        ("目标1", "梳理童话情节，简洁复述"),
        ("目标2", "赏析夸张、反讽等手法"),
        ("目标3", "探究「真话为何难以说出口」"),
    ], accent, "课时目标")

    add_question_slide(
        prs, bg,
        "爱新衣的皇帝、织空布的骗子、\n不敢说真话的大人、说真话的孩子——\n闹剧背后藏着什么？",
        "课前热身：真实与虚假的较量",
        M["coral"],
    )

    add_content_slide(prs, bg, "学习任务一 · 画流程图", [
        "每读完一段，记下「谁做了什么」",
        "用箭头串联骗局全过程",
        "骗子谎言→＿＿→＿＿→孩子说真话",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "学习任务二 · 拆心理机关", [
        "勾画老大臣、官员、皇帝的心理",
        "品味夸张语句的讽刺效果",
        "完成人物分析卡（证据—特点—作用）",
    ], M["mint"], "细读品味")

    add_question_slide(
        prs, bg,
        "人们都不敢说自己看不见，\n这是为什么？",
        "思考：他们共同害怕失去什么？",
        accent,
    )

    add_content_slide(prs, bg, "学习任务三 · 学会判断", [
        "完成课后「思考·探究·积累」",
        "联系生活：你有没有「随大流」的经历？",
        "主题二：面对虚假，要＿＿，不做沉默的大多数",
    ], M["sun"], "主题探究")

    add_summary_slide(prs, bg, [
        "夸张、重复、反差制造笑声",
        "笑声指向盲从与虚伪",
        "孩子的天真是不被虚假束缚的诚实",
    ], accent)

    add_content_slide(prs, bg, "作业", [
        ("必做", "150字概括故事"),
        ("情境", "班级群「震惊」消息，你会怎么做？（50字）"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第22课 皇帝的新装.pptx")


def ppt_lesson23():
    bg = ensure_bg("bg_nuwa.png", M["sage"], "nuwa")
    accent = M["sage"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "女娲造人", "第23课 · 第六、七课时 · 创造与生命", accent)

    add_content_slide(prs, bg, "本课目标", [
        ("目标1", "概括女娲造人过程，梳理思路"),
        ("目标2", "比较古籍与课文，体会想象创造"),
        ("目标3", "感悟「创造与生命」母题"),
    ], accent, "课时目标")

    add_question_slide(
        prs, bg,
        "为什么古人要想象\n一位女神来创造人类？",
        "这是对「人从哪里来、为何需要同伴」的诗意追问",
        M["lav"],
    )

    add_content_slide(prs, bg, "学习任务一 · 追踪创造", [
        "给女娲行动排序：孤独→＿＿→＿＿→繁衍",
        "比较两种造人方法的不同",
        "复述：从孤独到热闹的过程",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "学习任务二 · 比较阅读", [
        "阅读《风俗通》记载，完成比较表",
        "细读池水照影、黄泥捏人、藤条挥洒",
        "说说想象如何具体、生动、温暖",
    ], M["mint"], "细读品味")

    add_content_slide(prs, bg, "学习任务三 · 提炼策展主题", [
        "讨论：神话用想象解释人的来处",
        "女娲既有神性，也有人的孤独与喜悦",
        "主题三：生命来之不易，创造＿＿＿＿",
    ], M["coral"], "主题探究")

    add_summary_slide(prs, bg, [
        "神话改写基于古籍又超越古籍",
        "赋予创世故事现代意识与生命温度",
        "策展主题：创造与生命",
    ], accent)

    add_content_slide(prs, bg, "作业", [
        ("必做", "以「我看见女娲……」写100字画面"),
        ("选做", "查找一个中国创世神话比较"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第23课 女娲造人.pptx")


def ppt_lesson24():
    bg = ensure_bg("bg_fable.png", M["sun"], "fable")
    accent = M["sun"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "寓言四则", "第24课 · 第八、九课时 · 智慧与局限", accent)

    add_content_slide(prs, bg, "本课目标", [
        ("目标1", "归纳寓言文体特点，理解寓意"),
        ("目标2", "疏通文言，积累重点词语"),
        ("目标3", "探究寓意与情节设计，尝试新编"),
    ], accent, "课时目标")

    add_question_slide(
        prs, bg,
        "寓言是「穿着外套的真理」\n跨越时空，智慧相通",
        "课前热身：伊索、吕不韦、列子……",
        M["lav"],
    )

    add_content_slide(prs, bg, "学习任务一 · 初识寓言", [
        "默读四则，完成「寓言档案卡」",
        "讨论：蚊子平静离开，寓意会变吗？",
        "归纳寓言文体特点",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "学习任务二 · 深读寓意", [
        "解释加点词：闻、传、道、亡、只使……",
        "翻译：「得一人之使，非得一人于井中也」",
        "讨论「杞人忧天」的两种理解",
    ], M["mint"], "文言积累")

    add_content_slide(prs, bg, "学习任务三 · 寓言新编", [
        "任选一则，重新设计情节（100—200字）",
        "主题四：认识认知局限，＿＿＿＿",
        "同伴互评：寓意是否由情节推出？",
    ], M["coral"], "创意表达")

    add_summary_slide(prs, bg, [
        "寓言道理藏在欲望、选择与后果中",
        "读懂「为什么这样结尾」",
        "策展主题：智慧与局限",
    ], accent)

    add_content_slide(prs, bg, "作业", [
        ("必做", "《蚊子和狮子》现实启示"),
        ("背诵", "得一人之使，非得一人于井中也"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第24课 寓言四则.pptx")


def ppt_final_showcase():
    bg = ensure_bg("bg_final.png", M["lilac"], "final")
    accent = M["lilac"]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs, bg, "想象力博物馆开幕", "第十—十二课时 · 综合·写作·展示", accent)

    add_content_slide(prs, bg, "第十课时 · 制作说明卡", [
        "比较四类文本想象的依据与表现",
        "完成一张有文本依据的「想象说明卡」",
        "同伴审稿：事实准确？神奇具体？",
    ], M["sky"], "综合研读")

    add_content_slide(prs, bg, "想象的四种面貌", [
        ("神魔小说", "变化斗法，人情世故"),
        ("童话", "夸张反讽，照见人性"),
        ("神话", "创世温情，生命礼赞"),
        ("寓言", "简短警醒，寓意深远"),
    ], accent, "比较阅读")

    add_content_slide(prs, bg, "第十一课时 · 联想与想象", [
        "朗读《天上的街市》，辨析联想与想象",
        "三项要求：有依据、合逻辑、有新意",
        "完成创作提纲，写作不少于300字",
    ], M["mint"], "写作训练")

    add_question_slide(
        prs, bg,
        "联想是搭桥，想象是飞翔\n你的作品从哪里起飞？",
        "创作方向：续写童话 / 改写寓言 / 十年后的我 / 课本剧",
        M["coral"],
    )

    add_content_slide(prs, bg, "第十二课时 · 成果展示", [
        "布展：说明卡 + 创作作品",
        "讲解：30秒依据 + 60秒亮点 + 30秒思考",
        "参观者投「发现卡」并写理由",
    ], accent, "以生为主")

    add_content_slide(prs, bg, "回望母题链", [
        "自由 → 真实 → 创造 → 智慧",
        "成长就是在四者间寻找平衡",
        "填写自评表，整理学习档案",
    ], M["sun"], "单元总结")

    add_summary_slide(prs, bg, [
        "想象既是阅读能力，也是表达能力",
        "尊重文本和生活的逻辑",
        "看见习以为常之外的世界",
    ], accent)

    add_content_slide(prs, bg, "单元作业", [
        ("阅读", "完善策展卡，600字想象主题作文"),
        ("或", "5分钟以内课本剧脚本"),
        ("要求", "有依据、合逻辑、有新意"),
    ], M["peach"], "课后延伸")

    save_ppt(prs, "第10-12课时 单元成果展示.pptx")


def main():
    for fn in [
        ppt_unit_intro,
        ppt_lesson21,
        ppt_lesson22,
        ppt_lesson23,
        ppt_lesson24,
        ppt_final_showcase,
    ]:
        fn()
    print("\n全部生成完成！")


if __name__ == "__main__":
    main()
