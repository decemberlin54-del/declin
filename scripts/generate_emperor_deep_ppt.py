#!/usr/bin/env python3
"""Generate 9-slide deep teaching PPT for 《皇帝的新装》 — Morandi style."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/皇帝的新装深度教学PPT.pptx")
OUT_EN = Path("/workspace/emperor-new-clothes-deep-teaching.pptx")

# 1280x720 px @ 96dpi
SW = Inches(13.333)
SH = Inches(7.5)
PAD_L = Inches(60 / 96)
PAD_T = Inches(40 / 96)
PAD_R = Inches(60 / 96)
PAD_B = Inches(40 / 96)
INNER_W = SW - PAD_L - PAD_R
INNER_H = SH - PAD_T - PAD_B

BG = RGBColor(0xF4, 0xF1, 0xEA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
PRIMARY = RGBColor(0x6B, 0x90, 0x80)
SECONDARY = RGBColor(0xB5, 0x83, 0x8D)
ACCENT = RGBColor(0xE0, 0x9F, 0x68)
TEXT = RGBColor(0x2B, 0x2D, 0x42)
MUTED = RGBColor(0x84, 0xA9, 0x8C)


def rfont(run, size=20, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    sp = slide.shapes._spTree
    sp.remove(bg._element)
    sp.insert(2, bg._element)


def textbox(slide, left, top, width, height, lines, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, size, bold, color = item
            r = p.add_run()
            rfont(r, size, bold, color)
            r.text = text
        else:
            r = p.add_run()
            rfont(r)
            r.text = item
        p.space_after = Pt(6)
    return tb


def card(slide, left, top, width, height, border_color=PRIMARY):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = border_color
    sh.line.width = Pt(1.5)
    return sh


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    card(slide, PAD_L + Inches(0.3), PAD_T + Inches(0.5), INNER_W - Inches(0.6), INNER_H - Inches(1.0), PRIMARY)
    textbox(
        slide,
        PAD_L + Inches(0.8),
        PAD_T + Inches(1.2),
        INNER_W - Inches(1.6),
        Inches(1.2),
        [
            ("皇帝的新装", 40, True, PRIMARY),
            ("现实的讽刺与人性的镜子", 28, False, TEXT),
        ],
    )
    textbox(
        slide,
        PAD_L + Inches(0.8),
        PAD_T + Inches(2.6),
        INNER_W - Inches(1.6),
        Inches(0.5),
        [("统编版七年级上册第六单元第22课", 18, False, MUTED)],
    )
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PAD_L + Inches(0.8),
        PAD_T + Inches(3.4),
        Inches(4.2),
        Inches(1.6),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(0xE8, 0xE5, 0xDA)
    badge.line.color.rgb = SECONDARY
    textbox(
        slide,
        PAD_L + Inches(1.0),
        PAD_T + Inches(3.55),
        Inches(3.8),
        Inches(1.3),
        [
            ("［丹麦］安徒生", 20, True, SECONDARY),
            ("世界儿童文学巨匠", 16, False, TEXT),
            ("以夸张想象揭示人性", 16, False, MUTED),
        ],
    )
    textbox(
        slide,
        PAD_L + Inches(5.5),
        PAD_T + Inches(3.6),
        Inches(6.5),
        Inches(1.2),
        [
            ("二级母题：真实与虚假", 22, True, ACCENT),
            ("想象与真实 · 想象力博物馆", 18, False, TEXT),
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_goals(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("单元情境与学习目标", 32, True, PRIMARY)])
    cw = (INNER_W - Inches(0.3)) / 2
    card(slide, PAD_L, PAD_T + Inches(0.9), cw, Inches(4.8), PRIMARY)
    card(slide, PAD_L + cw + Inches(0.3), PAD_T + Inches(0.9), cw, Inches(4.8), ACCENT)
    textbox(
        slide,
        PAD_L + Inches(0.25),
        PAD_T + Inches(1.1),
        cw - Inches(0.5),
        Inches(4.4),
        [
            ("【单元情境】", 24, True, PRIMARY),
            ("", 8, False, TEXT),
            ("探索“想象与现实”的交融世界。", 20, False, TEXT),
            ("", 8, False, TEXT),
            ("本单元以“想象力博物馆”为情境，", 18, False, TEXT),
            ("在荒诞童话中照见真实，", 18, False, TEXT),
            ("在虚构故事里审视人心。", 18, False, TEXT),
        ],
    )
    textbox(
        slide,
        PAD_L + cw + Inches(0.55),
        PAD_T + Inches(1.1),
        cw - Inches(0.5),
        Inches(4.4),
        [
            ("【本课目标】", 24, True, ACCENT),
            ("", 8, False, TEXT),
            ("1. 理清“受骗—展骗—穿骗—看骗—揭骗”情节线索。", 18, False, TEXT),
            ("", 6, False, TEXT),
            ("2. 分析皇帝、大臣、百姓及孩子心理，", 18, False, TEXT),
            ("   剖析童话讽刺意味。", 18, False, TEXT),
            ("", 6, False, TEXT),
            ("3. 理解“虚伪与真实”的现实对应。", 18, False, TEXT),
        ],
    )


def slide_plot(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("故事推演脉络图", 32, True, PRIMARY)])
    steps = ["骗子做衣", "大臣探访", "皇帝试穿", "游行大典", "小孩揭穿"]
    sw = Inches(2.0)
    gap = Inches(0.35)
    start_x = PAD_L + Inches(0.2)
    y = PAD_T + Inches(1.2)
    for i, step in enumerate(steps):
        x = start_x + i * (sw + gap)
        card(slide, x, y, sw, Inches(1.4), PRIMARY if i < 4 else SECONDARY)
        textbox(slide, x + Inches(0.1), y + Inches(0.35), sw - Inches(0.2), Inches(0.8), [(step, 18, True, TEXT)])
        if i < 4:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw + Inches(0.02), y + Inches(0.55), gap - Inches(0.04), Inches(0.35))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()
    card(slide, PAD_L, PAD_T + Inches(3.0), INNER_W, Inches(2.2), SECONDARY)
    textbox(
        slide,
        PAD_L + Inches(0.4),
        PAD_T + Inches(3.2),
        INNER_W - Inches(0.8),
        Inches(1.8),
        [
            ("核心矛盾", 22, True, SECONDARY),
            ("虚荣心  vs  真实事实", 28, True, TEXT),
            ("谎言因“不敢承认看不见”而运转，真话因“孩子无畏”而刺破。", 18, False, MUTED),
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_characters(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("人物图鉴与心理剖析", 32, True, PRIMARY)])
    cols = [
        ("皇帝", "极致虚荣、愚蠢自私", "怕被认为不聪明、不称职", PRIMARY),
        ("大臣们", "自保阿谀、盲从从众", "怕失去官位与恩宠", ACCENT),
        ("骗子", "狡黠敏锐、利用弱点", "精准猎杀人性虚荣", SECONDARY),
    ]
    cw = (INNER_W - Inches(0.4)) / 3
    for i, (name, trait, psyche, color) in enumerate(cols):
        x = PAD_L + i * (cw + Inches(0.2))
        y = PAD_T + Inches(0.9)
        card(slide, x, y, cw, Inches(4.8), color)
        textbox(
            slide,
            x + Inches(0.2),
            y + Inches(0.25),
            cw - Inches(0.4),
            Inches(4.3),
            [
                (name, 26, True, color),
                ("", 8, False, TEXT),
                ("性格：" + trait, 18, False, TEXT),
                ("", 8, False, TEXT),
                ("心理：" + psyche, 18, False, TEXT),
            ],
        )


def slide_prop(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("关键道具与意象：“新装”", 32, True, PRIMARY)])
    textbox(
        slide,
        PAD_L,
        PAD_T + Inches(0.65),
        INNER_W,
        Inches(0.5),
        [("不存在的衣服，何以成为人性的“试金石”？", 20, False, SECONDARY)],
    )
    cw = (INNER_W - Inches(0.3)) / 2
    card(slide, PAD_L, PAD_T + Inches(1.3), cw, Inches(4.2), PRIMARY)
    card(slide, PAD_L + cw + Inches(0.3), PAD_T + Inches(1.3), cw, Inches(4.2), SECONDARY)
    textbox(
        slide,
        PAD_L + Inches(0.25),
        PAD_T + Inches(1.5),
        cw - Inches(0.5),
        Inches(3.8),
        [
            ("【新装的特征】", 22, True, PRIMARY),
            ("", 6, False, TEXT),
            ("· 愚蠢或不称职的人看不见", 18, False, TEXT),
            ("· 检验“聪明”与“称职”", 18, False, TEXT),
            ("· 以“看不见”制造恐惧", 18, False, TEXT),
        ],
    )
    textbox(
        slide,
        PAD_L + cw + Inches(0.55),
        PAD_T + Inches(1.5),
        cw - Inches(0.5),
        Inches(3.8),
        [
            ("【新装的实质】", 22, True, SECONDARY),
            ("", 6, False, TEXT),
            ("· 权力压迫下的群体谎言", 18, False, TEXT),
            ("· 虚荣与怯懦的遮羞布", 18, False, TEXT),
            ("· 沉默共谋的符号", 18, False, TEXT),
        ],
    )


def slide_child(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("为什么只有小孩敢说真话？", 32, True, PRIMARY)])
    cw = (INNER_W - Inches(0.3)) / 2
    card(slide, PAD_L, PAD_T + Inches(0.9), cw, Inches(4.5), SECONDARY)
    card(slide, PAD_L + cw + Inches(0.3), PAD_T + Inches(0.9), cw, Inches(4.5), PRIMARY)
    textbox(
        slide,
        PAD_L + Inches(0.25),
        PAD_T + Inches(1.1),
        cw - Inches(0.5),
        Inches(4.1),
        [
            ("成人的世界", 24, True, SECONDARY),
            ("", 6, False, TEXT),
            ("· 利益捆绑", 18, False, TEXT),
            ("· 社会化恐惧", 18, False, TEXT),
            ("· 虚伪掩面", 18, False, TEXT),
        ],
    )
    textbox(
        slide,
        PAD_L + cw + Inches(0.55),
        PAD_T + Inches(1.1),
        cw - Inches(0.5),
        Inches(4.1),
        [
            ("孩子的世界", 24, True, PRIMARY),
            ("", 6, False, TEXT),
            ("· 无利益纠葛", 18, False, TEXT),
            ("· 保持直觉", 18, False, TEXT),
            ("· 纯真赤诚", 18, False, TEXT),
        ],
    )
    hl = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PAD_L + Inches(2.5), PAD_T + Inches(5.7), Inches(5.5), Inches(0.7))
    hl.fill.solid()
    hl.fill.fore_color.rgb = RGBColor(0xF0, 0xE0, 0xE4)
    hl.line.fill.background()
    textbox(
        slide,
        PAD_L + Inches(2.7),
        PAD_T + Inches(5.75),
        Inches(5.1),
        Inches(0.6),
        [("重点：无私者无畏", 22, True, SECONDARY)],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_mirror(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("现实的镜子——童话中的现实批判", 30, True, PRIMARY)])
    blocks = [
        ("1. 官场生态", "上行下效的谄媚文化", PRIMARY),
        ("2. 社会心理", "群体无意识与沉默的螺旋", ACCENT),
        ("3. 权力困境", "失去真实反馈的上位者", SECONDARY),
    ]
    bh = Inches(1.5)
    gap = Inches(0.25)
    for i, (title, desc, color) in enumerate(blocks):
        y = PAD_T + Inches(0.95) + i * (bh + gap)
        card(slide, PAD_L, y, INNER_W, bh, color)
        textbox(
            slide,
            PAD_L + Inches(0.35),
            y + Inches(0.25),
            INNER_W - Inches(0.7),
            bh - Inches(0.4),
            [(title, 22, True, color), (desc, 20, False, TEXT)],
            anchor=MSO_ANCHOR.MIDDLE,
        )


def slide_mechanism(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("想象力的构建机制", 32, True, PRIMARY)])
    flow = [
        ("现实基础", "社会上的虚伪现象与谄媚风气", PRIMARY),
        ("夸张想象", "将“虚伪”实体化为“不存在的衣服”", ACCENT),
        ("逻辑闭环", "只要人人害怕承认，谎言就能完美运转", SECONDARY),
    ]
    bw = Inches(3.6)
    gap = Inches(0.45)
    y = PAD_T + Inches(1.5)
    for i, (title, desc, color) in enumerate(flow):
        x = PAD_L + i * (bw + gap)
        card(slide, x, y, bw, Inches(2.8), color)
        textbox(
            slide,
            x + Inches(0.2),
            y + Inches(0.35),
            bw - Inches(0.4),
            Inches(2.1),
            [(title, 22, True, color), ("", 6, False, TEXT), (desc, 18, False, TEXT)],
        )
        if i < 2:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.05), y + Inches(1.15), gap - Inches(0.1), Inches(0.4))
            ar.fill.solid()
            ar.fill.fore_color.rgb = MUTED
            ar.line.fill.background()
    card(slide, PAD_L, PAD_T + Inches(4.6), INNER_W, Inches(1.3), PRIMARY)
    textbox(
        slide,
        PAD_L + Inches(0.3),
        PAD_T + Inches(4.75),
        INNER_W - Inches(0.6),
        Inches(1.0),
        [("想象源于现实，夸张服务于讽刺，逻辑闭环维系谎言。", 20, False, TEXT)],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_task(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    textbox(slide, PAD_L, PAD_T, INNER_W, Inches(0.6), [("课堂拓展与写作任务", 32, True, PRIMARY)])
    card(slide, PAD_L, PAD_T + Inches(0.85), INNER_W, Inches(5.0), ACCENT)
    textbox(
        slide,
        PAD_L + Inches(0.4),
        PAD_T + Inches(1.1),
        INNER_W - Inches(0.8),
        Inches(4.5),
        [
            ("【任务】", 24, True, ACCENT),
            ("为《皇帝的新装》续写一个 200 字结尾。", 20, False, TEXT),
            ("", 8, False, TEXT),
            ("【提示】", 22, True, PRIMARY),
            ("游行大典结束后，皇帝回到宫中会发生什么？", 18, False, TEXT),
            ("大臣们会怎么做？百姓又会如何议论？", 18, False, TEXT),
            ("", 8, False, TEXT),
            ("【要求】", 22, True, SECONDARY),
            ("符合人物性格逻辑，具有讽刺效果。", 18, False, TEXT),
        ],
    )


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    slide_cover(prs)
    slide_goals(prs)
    slide_plot(prs)
    slide_characters(prs)
    slide_prop(prs)
    slide_child(prs)
    slide_mirror(prs)
    slide_mechanism(prs)
    slide_task(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    prs.save(str(OUT_EN))
    print(f"✓ {OUT} ({len(prs.slides)} slides)")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    build()
