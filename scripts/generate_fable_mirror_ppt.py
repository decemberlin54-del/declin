#!/usr/bin/env python3
"""Generate 8-slide Morandi PPT for 《寓言四则》短故事，大镜子."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/寓言四则短故事大镜子.pptx")
OUT_EN = Path("/workspace/fable-short-story-big-mirror.pptx")

SW = Inches(13.333)
SH = Inches(7.5)
PL, PT, PR, PB = Inches(0.625), Inches(0.42), Inches(0.625), Inches(0.42)
IW = SW - PL - PR

BG = RGBColor(0xF5, 0xF0, 0xE8)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x2C, 0x3E, 0x6B)
VERM = RGBColor(0xC4, 0x5C, 0x4A)
TEXT = RGBColor(0x2F, 0x3E, 0x46)
MUTED = RGBColor(0x6B, 0x7B, 0x8C)


def font(run, size=18, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    sh.fill.solid()
    sh.fill.fore_color.rgb = BG
    sh.line.fill.background()
    tree = slide.shapes._spTree
    tree.remove(sh._element)
    tree.insert(2, sh._element)


def tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(item, tuple):
            txt, sz, b, c = item
            r = p.add_run()
            font(r, sz, b, c)
            r.text = txt
        else:
            r = p.add_run()
            font(r)
            r.text = item
        p.space_after = Pt(4)
    return box


def rounded(slide, l, t, w, h, color=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = color
    sh.line.width = Pt(1.5)
    return sh


def header(slide, title, num):
    tb(slide, PL, PT, IW * 0.7, Inches(0.35), [(title, 22, True, NAVY)])
    tb(slide, PL + IW * 0.75, PT, IW * 0.25, Inches(0.35), [(f"{num:02d} / 08", 14, False, MUTED)], PP_ALIGN.RIGHT)
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PL, PT + Inches(0.38), IW, Inches(0.03))
    ln.fill.solid()
    ln.fill.fore_color.rgb = VERM
    ln.line.fill.background()


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    mirror = rounded(s, PL + IW * 0.28, PT + Inches(0.9), IW * 0.44, Inches(3.6), VERM)
    tb(s, PL, PT + Inches(0.2), IW, Inches(0.5), [("《寓言四则》——短故事，大镜子", 36, True, NAVY)], PP_ALIGN.CENTER)
    tb(s, PL, PT + Inches(0.75), IW, Inches(0.4), [("课时六、七  教学演示文稿", 20, False, MUTED)], PP_ALIGN.CENTER)
    for i, (txt, x, y) in enumerate([
        ("赫耳墨斯\n和雕像者", 0.05, 1.1),
        ("蚊子\n和狮子", 0.72, 1.1),
        ("穿井\n得一人", 0.05, 3.5),
        ("杞人\n忧天", 0.72, 3.5),
    ]):
        rounded(s, PL + IW * x, PT + Inches(y), IW * 0.22, Inches(0.9), NAVY)
        tb(s, PL + IW * x + Inches(0.05), PT + Inches(y + 0.12), IW * 0.2, Inches(0.7), [(txt, 13, True, NAVY)], PP_ALIGN.CENTER)
    tb(s, PL, SH - Inches(0.7), IW, Inches(0.4), [("短故事照见人心 · 大镜子折射现实", 16, False, VERM)], PP_ALIGN.CENTER)


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "情境导入与课时目标", 2)
    rounded(s, PL, PT + Inches(0.55), IW, Inches(1.0), VERM)
    tb(s, PL + Inches(0.15), PT + Inches(0.65), IW - Inches(0.3), Inches(0.9), [
        "为什么《赫耳墨斯》中雕像者的一句话会让他难堪？",
        "为什么《穿井得一人》的一句话会在国都传开？",
    ])
    cw = (IW - Inches(0.2)) / 2
    rounded(s, PL, PT + Inches(1.75), cw, Inches(2.5), NAVY)
    tb(s, PL + Inches(0.12), PT + Inches(1.85), cw - Inches(0.24), Inches(2.3), [
        ("【课时目标】", 16, True, NAVY),
        "1. 归纳寓言文体特点，理解各则寓意",
        "2. 疏通文言，积累重点词语",
        "3. 分析情节与寓意关系，探究多种理解",
    ])
    rounded(s, PL + cw + Inches(0.2), PT + Inches(1.75), cw, Inches(2.5), VERM)
    tb(s, PL + cw + Inches(0.32), PT + Inches(1.85), cw - Inches(0.24), Inches(2.3), [
        ("【评价任务】", 16, True, VERM),
        "① 完成寓言档案卡（目标1）",
        "② 翻译重点句，解释文言词（目标2）",
        "③ 分析寓意，联系生活谈启示（目标3）",
    ])


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "学习任务一 · 寓言档案卡", 3)
    rows = [
        ["篇目", "人物", "情节", "寓意"],
        ["赫耳墨斯", "赫耳墨斯、雕像者", "问价自取其辱", "讽刺妄自尊大"],
        ["蚊子和狮子", "蚊子、狮子", "胜后骄纵被蛛网困", "骄兵必败"],
        ["穿井得一人", "丁氏、宋君", "传言失真宋君求证", "以讹传讹须求证"],
        ["杞人忧天", "杞人、晓之者", "忧天被开导释然", "不必杞人忧天"],
    ]
    y = PT + Inches(0.55)
    rh = Inches(0.42)
    cw = [IW * 0.18, IW * 0.22, IW * 0.32, IW * 0.28]
    x = PL
    for ri, row in enumerate(rows):
        x = PL
        for ci, cell in enumerate(row):
            rounded(s, x, y, cw[ci] - Inches(0.04), rh, NAVY if ri == 0 else MUTED)
            tb(s, x + Inches(0.05), y + Inches(0.06), cw[ci] - Inches(0.1), rh - Inches(0.08),
               [(cell, 12 if ri else 11, ri == 0, CARD if ri == 0 else TEXT)])
            x += cw[ci]
        y += rh + Inches(0.05)
    rounded(s, PL, y + Inches(0.1), IW, Inches(1.1), VERM)
    tb(s, PL + Inches(0.12), y + Inches(0.18), IW - Inches(0.24), Inches(0.95), [
        ("寓言文体特点", 14, True, VERM),
        "篇幅：短小精悍  |  人物：典型鲜明  |  写法：拟人、夸张、反转  |  作用：寄寓道理、警醒世人",
    ])


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "学习任务二 · 品味情节，读懂寓意", 4)
    cw = (IW - Inches(0.15)) / 2
    rounded(s, PL, PT + Inches(0.55), cw, Inches(2.0), NAVY)
    tb(s, PL + Inches(0.1), PT + Inches(0.65), cw - Inches(0.2), Inches(1.8), [
        ("情节 → 寓意", 15, True, NAVY),
        "赫耳墨斯：「白送」→ 讽刺自高自大",
        "蚊子：凯歌后被粘 → 骄兵必败",
        "丁氏：澄清误会 → 勿以讹传讹",
    ])
    rounded(s, PL + cw + Inches(0.15), PT + Inches(0.55), cw, Inches(2.0), VERM)
    tb(s, PL + cw + Inches(0.25), PT + Inches(0.65), cw - Inches(0.2), Inches(1.8), [
        ("改写探究", 15, True, VERM),
        "若蚊子战胜后悄悄离开……",
        "寓意或变为「懂得适可而止」",
        "说明：寓意与情节设计密切相关",
    ])
    rounded(s, PL, PT + Inches(2.7), IW, Inches(1.5), NAVY)
    tb(s, PL + Inches(0.12), PT + Inches(2.8), IW - Inches(0.24), Inches(1.3), [
        ("联系生活", 14, True, NAVY),
        "① 未经核实转发消息 → 《穿井得一人》  ② 成绩好看不起同学 → 《赫耳墨斯》  ③ 风声就焦虑 → 《杞人忧天》",
    ])


def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "学习任务三 · 疏通文言，积累词语", 5)
    words = [("闻", "听说"), ("道", "讲述"), ("亡", "同「无」"), ("晓", "开导"), ("舍然", "释然")]
    x = PL
    for w, m in words:
        rounded(s, x, PT + Inches(0.55), Inches(2.2), Inches(0.85), NAVY)
        tb(s, x + Inches(0.08), PT + Inches(0.62), Inches(2.0), Inches(0.7), [(f"{w}：{m}", 14, True, NAVY)])
        x += Inches(2.35)
    rounded(s, PL, PT + Inches(1.6), IW, Inches(2.6), VERM)
    tb(s, PL + Inches(0.12), PT + Inches(1.72), IW - Inches(0.24), Inches(2.4), [
        ("重点句翻译", 15, True, VERM),
        "（1）得一人之使，非得一人于井中也。",
        "（2）求闻之若此，不若无闻也。",
        "（3）若屈伸呼吸，终日在天中行止，奈何忧崩坠乎？",
    ])


def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "「杞人忧天」的多视角解读", 6)
    cw = (IW - Inches(0.15)) / 2
    rounded(s, PL, PT + Inches(0.55), cw, Inches(2.8), NAVY)
    tb(s, PL + Inches(0.1), PT + Inches(0.65), cw - Inches(0.2), Inches(2.6), [
        ("视角 A：讽刺不必要的杞忧", 15, True, NAVY),
        "依据：「废寝食者」「奈何忧崩坠乎」",
        "启示：遇事应实事求是，勿凭空忧虑",
    ])
    rounded(s, PL + cw + Inches(0.15), PT + Inches(0.55), cw, Inches(2.8), VERM)
    tb(s, PL + cw + Inches(0.25), PT + Inches(0.65), cw - Inches(0.2), Inches(2.6), [
        ("视角 B：体现忧患意识", 15, True, VERM),
        "依据：「忧天地崩坠，身亡所寄」",
        "启示：对未知保持审慎，亦需理性求证",
    ])


def slide7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "给古人一条今天的建议", 7)
    steps = ["丁氏原话", "传言失真", "国人传播", "宋君求证"]
    x = PL
    for i, st in enumerate(steps):
        rounded(s, x, PT + Inches(0.6), Inches(2.5), Inches(0.7), NAVY)
        tb(s, x + Inches(0.05), PT + Inches(0.7), Inches(2.4), Inches(0.5), [(st, 13, True, NAVY)], PP_ALIGN.CENTER)
        if i < 3:
            tb(s, x + Inches(2.5), PT + Inches(0.75), Inches(0.35), Inches(0.4), [("→", 18, True, VERM)])
        x += Inches(2.85)
    rounded(s, PL, PT + Inches(1.6), IW, Inches(1.0), VERM)
    tb(s, PL + Inches(0.12), PT + Inches(1.7), IW - Inches(0.24), Inches(0.85), [
        ("面对网络信息：核对来源 · 理性判断 · 不传谣言", 16, True, VERM),
    ])
    rounded(s, PL, PT + Inches(2.8), IW, Inches(1.2), NAVY)
    tb(s, PL + Inches(0.12), PT + Inches(2.9), IW - Inches(0.24), Inches(1.0), [
        ("课堂小结", 14, True, NAVY),
        "两则文言寓言警示我们：信息须求证，忧虑须理性，智慧照见局限。",
    ])


def slide8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, "作业与检测", 8)
    cw = (IW - Inches(0.15)) / 2
    rounded(s, PL, PT + Inches(0.55), cw, Inches(2.8), NAVY)
    tb(s, PL + Inches(0.1), PT + Inches(0.65), cw - Inches(0.2), Inches(2.6), [
        ("必做", 16, True, NAVY),
        "1. 用「这则寓言告诉我们……」写《蚊子和狮子》启示（80字）",
        "2. 背诵：得一人之使，非得一人于井中也",
    ])
    rounded(s, PL + cw + Inches(0.15), PT + Inches(0.55), cw, Inches(2.8), VERM)
    tb(s, PL + cw + Inches(0.25), PT + Inches(0.65), cw - Inches(0.2), Inches(2.6), [
        ("延伸拓展", 16, True, VERM),
        "选做：任选一则寓言新编（100—200字）",
        "拓展：搜集一则网络谣言案例，与《穿井得一人》对比分析",
    ])


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    for fn in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    import shutil
    shutil.copy2(OUT, OUT_EN)
    print(f"✓ {OUT}")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    main()
