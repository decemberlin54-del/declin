#!/usr/bin/env python3
"""Generate 10-slide Morandi PPT for Unit 6 Lesson 1-2."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path("/workspace/第六单元PPT/第1-2课时神技背后的真实10页PPT.pptx")
OUT_EN = Path("/workspace/unit6-lesson1-2-ten-slides.pptx")

SW, SH = Inches(13.333), Inches(7.5)
PL, PT, PR, PB = Inches(60 / 96), Inches(40 / 96), Inches(60 / 96), Inches(40 / 96)
IW, IH = SW - PL - PR, SH - PT - PB

BG = RGBColor(0xF5, 0xF3, 0xEC)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x6B, 0x90, 0x80)
BLUE = RGBColor(0x58, 0x7B, 0x7C)
PINK = RGBColor(0xB5, 0x83, 0x8D)
CLAY = RGBColor(0xE0, 0x9F, 0x68)
TEXT = RGBColor(0x2F, 0x3E, 0x46)
MUTED = RGBColor(0x84, 0xA9, 0x8C)


def font(run, size=20, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bg(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    s.fill.solid()
    s.fill.fore_color.rgb = BG
    s.line.fill.background()
    sp = slide.shapes._spTree
    sp.remove(s._element)
    sp.insert(2, s._element)


def box(slide, l, t, w, h, lines, border=GREEN):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = border
    card.line.width = Pt(1.5)
    tf = slide.shapes.add_textbox(l + Inches(0.2), t + Inches(0.15), w - Inches(0.4), h - Inches(0.3)).text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        t2, sz, bd, c = item if len(item) == 4 else (item, 20, False, TEXT)
        r = p.add_run()
        font(r, sz, bd, c)
        r.text = t2
        p.space_after = Pt(4)


def title(slide, text, top=PT):
    tb = slide.shapes.add_textbox(PL, top, IW, Inches(0.7))
    r = tb.text_frame.paragraphs[0].add_run()
    font(r, 32, True, GREEN)
    r.text = text


def footer(slide, text):
    tb = slide.shapes.add_textbox(PL, SH - PB - Inches(0.55), IW, Inches(0.45))
    r = tb.text_frame.paragraphs[0].add_run()
    font(r, 18, False, PINK)
    r.text = text


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # 1 封面
    s = new_slide(prs)
    box(s, PL + Inches(1), PT + Inches(1.2), IW - Inches(2), Inches(4.5), [
        ("神技背后的真实", 40, True, TEXT),
        ("自由边界与生命创造", 28, False, GREEN),
        ("", 10, False, TEXT),
        ("七年级上 · 第六单元 · 第1—2课时", 18, False, MUTED),
        ("想象力博物馆 · 想象与真实", 20, False, CLAY),
    ], GREEN)
    footer(s, "在奇思妙想中照见真实，在神技变化中读懂人心。")

    # 2 单元情境
    s = new_slide(prs)
    title(s, "你敢挑战吗？")
    cw = (IW - Inches(0.25)) / 2
    box(s, PL, PT + Inches(0.9), cw, Inches(3.5), [
        ("【大任务】", 22, True, GREEN),
        ("举办「想象力博物馆」微展览", 18, False, TEXT),
        ("制作想象说明卡", 18, True, PINK),
    ], GREEN)
    box(s, PL + cw + Inches(0.25), PT + Inches(0.9), cw, Inches(3.5), [
        ("【母题链】", 22, True, CLAY),
        ("自由→真实→创造→智慧", 18, False, TEXT),
    ], CLAY)
    footer(s, "四种想象文体，一条「想象与真实」母题链。")

    # 3 核心双栏
    s = new_slide(prs)
    title(s, "神技背后的真实")
    box(s, PL, PT + Inches(0.85), cw, Inches(4.2), [
        ("孙悟空：自由与规则", 24, True, GREEN),
        ("变土地庙为何露出尾巴？", 20, False, BLUE),
        ("神幻想象的物理约束", 20, True, PINK),
        ("变形须遵循事物原型特征", 18, False, TEXT),
    ], GREEN)
    box(s, PL + cw + Inches(0.25), PT + Inches(0.85), cw, Inches(4.2), [
        ("女娲：创造与温情", 24, True, CLAY),
        ("为何从「手捏」变成「挥泥」？", 20, False, BLUE),
        ("情感具象化", 20, True, CLAY),
        ("折射抚育子女的辛劳与大爱", 18, False, TEXT),
    ], CLAY)
    footer(s, "真正的想象，是现实逻辑与人类情感在奇幻世界里的投射。")

    # 4 学习目标
    s = new_slide(prs)
    title(s, "你将学会什么？")
    box(s, PL, PT + Inches(0.9), IW, Inches(3.8), [
        ("· 快速默读：节点词+简洁句概括情节", 20, False, TEXT),
        ("· 比较神魔、神话想象的依据与特点", 20, False, TEXT),
        ("· 证据—特点—作用 分析人物", 20, True, PINK),
    ], BLUE)
    footer(s, "既读得进奇幻，也读得出真实。")

    # 5 斗法
    s = new_slide(prs)
    title(s, "七十二变：自由与克制")
    steps = ["麻雀", "鱼儿", "庙宇", "旗竿破绽"]
    sw = Inches(2.3)
    for i, st in enumerate(steps):
        x = PL + i * (sw + Inches(0.35))
        box(s, x, PT + Inches(1.0), sw, Inches(1.2), [(st, 18, True, TEXT)], GREEN)
    box(s, PL, PT + Inches(2.5), IW, Inches(1.5), [
        ("变化再奇，也受相克规则制约——神魔皆有人情", 20, False, TEXT),
    ], GREEN)
    footer(s, "策展主题一：自由可贵，但须守规则。")

    # 6 女娲
    s = new_slide(prs)
    title(s, "创世神话的温度")
    box(s, PL, PT + Inches(0.9), cw, Inches(2.8), [
        ("古籍：女娲抟黄土作人", 20, False, TEXT),
    ], CLAY)
    box(s, PL + cw + Inches(0.25), PT + Inches(0.9), cw, Inches(2.8), [
        ("课文：孤独、喜悦、挥藤沾泥", 20, False, TEXT),
        ("神性与人性融合", 20, True, CLAY),
    ], CLAY)
    footer(s, "策展主题二：创造与生命。")

    # 7 对比
    s = new_slide(prs)
    title(s, "两种想象，两种面向")
    y0 = PT + Inches(0.95)
    box(s, PL, y0, IW, Inches(2.8), [
        ("神魔小说：相克斗法 → 自由与规则", 20, False, TEXT),
        ("神话改写：添情创造 → 创造与生命", 20, False, TEXT),
        ("同有想象，各有不同表达效果", 18, False, MUTED),
    ], BLUE)
    footer(s, "比较阅读，发现规律。")

    # 8 方法
    s = new_slide(prs)
    title(s, "怎么读想象类文本？")
    box(s, PL, PT + Inches(0.9), IW, Inches(1.3), [("默读三要素：谁 — 做什么 — 结果怎样", 22, True, PINK)], BLUE)
    box(s, PL, PT + Inches(2.4), IW, Inches(1.3), [("赏析三步：证据 — 特点 — 作用", 22, True, PINK)], BLUE)
    footer(s, "追问：它像现实中的什么？")

    # 9 笔记
    s = new_slide(prs)
    title(s, "想象说明卡 · 初稿")
    box(s, PL, PT + Inches(0.9), IW, Inches(3.5), [
        ("文本片段 → 想象依据 → 表现方式 → 现实意味", 20, False, BLUE),
        ("为展品写清：从哪来、怎样写、有何用", 22, True, GREEN),
    ], GREEN)
    footer(s, "为「想象力博物馆」策展，从阅读笔记开始。")

    # 10 小结
    s = new_slide(prs)
    title(s, "今日收获")
    box(s, PL, PT + Inches(0.9), IW, Inches(3.2), [
        ("神魔想象重 自由与规则", 20, False, TEXT),
        ("神话想象重 创造与生命", 20, False, TEXT),
        ("作业：用100字写课文梗概", 20, True, CLAY),
    ], GREEN)
    footer(s, "下节继续：童话里的真实与虚假。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    prs.save(str(OUT_EN))
    print(f"✓ {OUT} ({len(prs.slides)} slides)")
    print(f"✓ {OUT_EN}")


if __name__ == "__main__":
    build()
